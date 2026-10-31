from datetime import date, timedelta
from collections import defaultdict
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook
from report_cleaner.case_manager import build_report, plan_fields
from report_cleaner.case_manager_sources import Sources
from report_cleaner.case_manager_rules import (
    reclassify, generic_name, radiation_review, death_review, new_courses,
    fraction_count, holiday_calendar, planned_drugs,
)
from report_cleaner.case_manager_excel import export_report, sync_manual
from test_case_manager import case, event


@pytest.fixture
def root(tmp_path):
    shutil.copytree(Path(__file__).parents[1]/'mapping',tmp_path/'mapping')
    return tmp_path


@pytest.mark.parametrize('code,name,drug,category,scope',[
    ('63010C','切片',False,'未分類醫令',''),
    ('67051B','肺手術',False,'手術','LC'),
    ('79417B','手術',False,'手術','PR'),
    ('36011B','照射',False,'放射治療',''),
    ('36015B','計畫',False,'放療計畫',''),
    ('OMET3','Methotrexate 2.5mg',True,'未分類藥物',''),
    ('TEST','Methotrexate 50mg injection',True,'化學治療',''),
    ('TEST','Methotrexate 2.5mg injection',True,'化學治療',''),
    ('TEST','Tegafur/Uracil',True,'化學治療',''),
    ('TEST','Flutamide',True,'抗荷爾蒙治療',''),
    ('TEST','Sotorasib',True,'標靶治療',''),
    ('TEST','Denosumab',True,'未分類藥物',''),
])
def test_reclassification(code,name,drug,category,scope):
    e=reclassify(event(code,category='化學治療' if drug else '未分類醫令',scope='',name=name,drug=drug))
    assert (e['category'],e['scope'])==(category,scope)


def test_generic_and_plan_drug_dictionary(root):
    assert generic_name('Gemcitabine HCl 200mg')=='Gemcitabine'
    assert generic_name('Gemcitabine 1000mg')=='Gemcitabine'
    assert planned_drugs({'化學治療':'EpiC','抗荷爾蒙治療':'Lenozole'},root/'mapping/drug_treatment_mapping.csv')==['Cyclophosphamide','Epirubicin','Letrozole']


def test_liver_plan_notes_not_radiation():
    p=case(cancer='LI_肝癌',放療其他說明='已完成TAE',RFA='N_否')
    fields=plan_fields([p])
    assert fields['放射治療']==''
    assert fields['肝動脈栓塞（TACE）']
    assert not fields['射頻燒灼（RFA）']


@pytest.mark.parametrize('reason',['轉院-院所','死亡(未治療)','拒絕治療'])
def test_class_zero_excludes_all_tasks(root,reason):
    s=Sources(cases=[case(個案分類='Class 0',結案原因=reason)],events=[event()],as_of=date(2026,10,20))
    r=build_report(s,root)
    assert not r.tasks
    assert '計畫書' not in r.planned[0]['資料提醒']
    assert r.actual[0]['是否符合計畫書'].startswith('不比對')


def test_two_way_comparison_and_supportive(root):
    s=Sources(cases=[case()],plans=[case(標靶治療方式='藥物')],events=[event(),event('SUP','化學治療','',drug=True,name='Leucovorin')],as_of=date(2026,10,20))
    r=build_report(s,root)
    assert r.actual[0]['是否符合計畫書'].startswith('否')
    assert {'計畫書缺漏','計畫書有、查無實際治療'}<={t['待辦類型'] for t in r.tasks}
    assert '未分類醫令／藥物' not in r.actual[0]['資料提醒']
    assert r.actual[0]['藥物種類數']==0


def categories(*events):
    result=defaultdict(list)
    for e in events: result[e['category']].append(e)
    return result


def rt(d): return event('36012B','放射治療','',date=d)


def test_spring_holiday_and_replan():
    c=categories(rt(date(2026,2,13)),rt(date(2026,2,23)))
    assert not radiation_review(c,[],[],date(2026,3,1),[])[3]
    c=categories(rt(date(2026,7,13)),rt(date(2026,7,21)),event('36015B','放療計畫','',date=date(2026,7,22)))
    result=radiation_review(c,[],[],date(2026,8,1),[])
    assert len(result[3])==1 and result[3][0]['replan']==date(2026,7,22)
    assert result[3][0]['days']==1
    assert {d.year for d in holiday_calendar()}==set(range(2019,2028))


def test_radiation_completion_all_segments_and_ccrt():
    days=[date(2026,1,2)+timedelta(days=n) for n in range(5)]+[date(2026,2,3)+timedelta(days=n) for n in range(5)]
    c=categories(*[rt(d) for d in days],event('CHEMO','化學治療','',date=date(2026,1,9)))
    p=case(放射線治療填寫日期='2026-01-01',DOSECGY='50 Gy x 20')
    completion,incomplete,_,_,ccrt=radiation_review(c,[p],[date(2026,1,1)],date(2026,3,1),[])
    assert '實際10次／計畫20次' in completion and incomplete
    assert '疑似同步' in ccrt
    assert fraction_count('15fx + boost 5fx')==15
    assert fraction_count('5000 cGy') is None


def test_new_course_strict_threshold_and_class_zero():
    start=date(2026,1,1)
    c=categories(event('CHEMO','化學治療','',date=start+timedelta(days=184)))
    assert len(new_courses(c,[start],False))==1
    assert not new_courses(c,[start],True)
    assert not new_courses(c,[start+timedelta(days=1)],False)


def test_death_per_modality_palliative_and_future_orders():
    death=date(2026,9,1)
    c=categories(event(date=date(2026,8,5)),event('CHEMO','化學治療','',date=date(2026,8,20)),rt(date(2026,9,2)))
    p=case(化學治療目的1='Palliative')
    result=death_review(death,'',c,[p])
    assert result.startswith('疑似必要提報：手術') and '已排除Palliative' in result
    assert '放射治療' not in result
    del c['手術']
    assert death_review(death,'',c,[p]).startswith('排除：治療目的為Palliative')
    p['化學治療目的2']='Curative'
    assert death_review(death,'',c,[p]).startswith('疑似必要提報')
    assert death_review(death,'死亡(未治療)',c,[p]).startswith('排除：結案')


def test_timing_visit_fallback_mdt_and_formulas(root):
    s=Sources(cases=[case(最近看診日='2027-01-01',最近掛號='2026-09-01',期別='III')],events=[event(date=date(2026,6,15)),event(date=date(2026,8,5))],as_of=date(2026,10,20))
    r=build_report(s,root)
    assert r.actual[0]['首次治療日']==date(2026,8,5)
    assert r.actual[0]['最後回診日']==date(2026,9,1)
    assert '建議提MDT' in r.actual[0]['MDT討論']
    assert {t['嚴重度'] for t in r.tasks}<={'紅','黃','灰'}
    m=root/'review/manual.xlsx'; records=sync_manual(r,m)
    out=root/'report.xlsx'; export_report(r,out,records,m)
    w=load_workbook(out)
    assert 'TODAY()' in w['病人實際治療總表']['R2'].value
    assert 'TODAY()' in w['病人實際治療總表']['S2'].value
    assert w['統計摘要']['B15'].data_type=='f'
    assert w['國定假日'].max_row>150
    assert '期別確認紀錄' in w.sheetnames
