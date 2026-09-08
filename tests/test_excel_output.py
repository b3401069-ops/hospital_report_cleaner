from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from report_cleaner.cleaner import (
    PATIENT_LIST_COLUMNS,
    PATIENT_MASTER_COLUMNS,
    export_cleaned_report,
    _build_patient_list_zh,
)


def _row(columns: list[str], **values: str) -> dict[str, str]:
    row = {column: "" for column in columns}
    row.update(values)
    return row


def test_patient_view_preserves_multi_case_summaries_and_unknown_treatments() -> None:
    summary = "癌別甲：臨床TNM／期別 II | 癌別乙：病理TNM／期別 III"
    drugs = " | ".join(f"測試藥物{i}" for i in range(35))
    source = pd.DataFrame([{
        "medical_record_no": "00123", "patient_name": "測試個案",
        "final_stage": "II | III", "tnm_summary": summary,
        "medication_summary": drugs, "chemotherapy": "Y",
    }])
    result = _build_patient_list_zh(source)
    assert len(result) == 1
    assert result.loc[0, "病歷號"] == "00123"
    assert result.loc[0, "最終期別"] == "II | III"
    assert result.loc[0, "TNM／期別摘要"] == summary
    assert result.loc[0, "藥物摘要"] == drugs
    assert result.loc[0, "化學治療"] == "有"
    assert result.loc[0, "標靶治療"] == ""
    assert source.loc[0, "chemotherapy"] == "Y"


def test_export_creates_chinese_first_page_and_readable_patient_table(
    tmp_path: Path,
) -> None:
    patient_master = pd.DataFrame(
        [
            _row(
                PATIENT_MASTER_COLUMNS,
                patient_name="測試個案",
                sex="女",
                medical_record_no="TEST001",
                age_at_diagnosis="65",
                diagnosis_date="2025-02-03",
                cancer_type="BR_乳癌",
                cancer_summary="乳癌（BR）",
                tnm_summary="乳癌（BR）：期別 IIA",
                diagnosis_text="乳房惡性腫瘤",
                final_stage="IIA",
                doctor="測試醫師",
                chemotherapy="Y",
                surgery="Y",
                patient_key="MRN:TEST001",
            )
        ]
    )
    patient_list = pd.DataFrame(
        [
            _row(
                PATIENT_LIST_COLUMNS,
                patient_name="測試個案",
                sex="女",
                medical_record_no="TEST001",
                age_at_diagnosis="65",
                diagnosis_date="2025-02-03",
                cancer_type="BR_乳癌",
                cancer_summary="乳癌（BR）",
                tnm_summary="乳癌（BR）：期別 IIA",
                diagnosis_text="乳房惡性腫瘤",
                final_stage="IIA",
                doctor="測試醫師",
                chemotherapy="Y",
                surgery="Y",
            )
        ]
    )
    standardized = patient_master.assign(
        report_name="測試來源",
        source_file="sample.xlsx",
        record_key="R1",
        order_category="drug",
        order_code="TESTDRUG",
        order_name="測試藥物",
    )
    empty = pd.DataFrame()
    stage_review = pd.DataFrame(
        [
            {
                "cancer_type": "LC_肺癌",
                "clinical_tnm": "cT3N1M0",
                "pathologic_tnm": "",
                "primary_tnm": "cT3N1M0",
                "final_stage": "",
                "final_stage_source": "",
                "review_reason": "missing_final_stage",
                "count": 1,
            }
        ]
    )

    output = export_cleaned_report(
        standardized,
        standardized,
        patient_master,
        patient_list,
        empty,
        stage_review,
        empty,
        empty,
        tmp_path / "cleaned_reports.xlsx",
    )

    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames[:7] == [
        "使用說明",
        "病人查詢",
        "病人總表",
        "統計摘要",
        "月趨勢",
        "治療藥物明細",
        "抽查清單",
    ]
    assert workbook.active.title == "病人查詢"
    assert workbook.calculation.calcMode == "auto"
    assert workbook["使用說明"]["A1"].value == "癌症報表清洗結果"
    assert not workbook["使用說明"].merged_cells.ranges
    assert workbook["使用說明"]["C1"].fill.fgColor.rgb.endswith("17365D")
    assert workbook["使用說明"]["B12"].value == 1
    assert workbook["使用說明"]["B13"].value == 1
    assert workbook["使用說明"]["B15"].value == 1

    patient_sheet = workbook["病人總表"]
    assert patient_sheet.freeze_panes == "C2"
    assert patient_sheet.auto_filter.ref.endswith(str(patient_sheet.max_row))
    headers = [cell.value for cell in patient_sheet[1]]
    assert headers[:10] == [
        "病患姓名",
        "病歷號",
        "性別",
        "年齡",
        "癌別摘要",
        "最終期別",
        "期別資料狀態",
        "最後回診日",
        "最後資料日期",
        "主責醫師",
    ]
    positions = {cell.value: cell.column for cell in patient_sheet[1]}
    assert patient_sheet.cell(2, positions["治療摘要"]).value == "化療、手術"
    assert "缺回診日" in patient_sheet.cell(2, positions["資料提醒"]).value
    assert patient_sheet.cell(2, positions["化學治療"]).value == "有"
    assert patient_sheet.cell(2, positions["手術"]).value == "有"
    assert patient_sheet.cell(2, positions["標靶治療"]).value is None
    assert patient_sheet.column_dimensions["V"].hidden
    assert "PatientOverview" in patient_sheet.tables
    assert "病人期別待確認" in workbook.sheetnames
    search = workbook["病人查詢"]
    assert search["B3"].number_format == "@"
    assert search["B3"].value is None
    assert "MATCH(" in search["H3"].value
    assert "COUNTIF(" in search["H2"].value
    assert "同病歷號有多列" in search["B4"].value
    assert "INDEX(" in search["B6"].value
    assert "HYPERLINK(" in search["B29"].value

    detail_sheet = workbook["治療藥物明細"]
    assert detail_sheet.freeze_panes == "E2"
    assert detail_sheet["G2"].value == "TESTDRUG"
    assert detail_sheet["H2"].value == "測試藥物"
    assert detail_sheet["I2"].value == 1

    qa_sheet = workbook["抽查清單"]
    assert qa_sheet.freeze_panes == "E2"
    assert qa_sheet["M1"].value == "人工確認結果"
    assert qa_sheet.data_validations.count == 1

    assert workbook["TNM期別待確認"]["G1"].value == "待確認原因"
    assert workbook["TNM期別待確認"]["G2"].value == "有 TNM，缺最終期別"
    for technical_sheet in (
        "cleaned_data",
        "standardized_data",
        "patient_master",
        "patient_list",
        "stage_review",
    ):
        assert workbook[technical_sheet].sheet_state == "hidden"
