from pathlib import Path
import shutil

import pandas as pd
import pytest
from openpyxl import load_workbook

from report_cleaner.patient_stage_review import apply_confirmations, build_queue, sync_review


def records() -> pd.DataFrame:
    return pd.DataFrame([
        {"patient_key": "MRN:001", "medical_record_no": "001", "case_id": "A",
         "patient_name": "測試甲", "cancer_type": "LC_肺癌", "diagnosis_date": "2025-01-01",
         "clinical_tnm": "cT2N0M0", "pathologic_tnm": "", "final_stage": "",
         "final_stage_source": "", "report_name": "TNM報表", "source_file": "sample.xlsx"},
        {"patient_key": "MRN:002", "medical_record_no": "002", "case_id": "B",
         "patient_name": "測試乙", "cancer_type": "LC_肺癌", "diagnosis_date": "2025-01-01",
         "clinical_tnm": "cT2N0M0", "pathologic_tnm": "", "final_stage": "",
         "final_stage_source": "", "report_name": "TNM報表", "source_file": "sample.xlsx"},
    ])


def approve(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["個案確認"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for label, value in {
        "確認期別": "II", "分期類別": "臨床", "確認依據": "測試用核對紀錄",
        "確認人": "測試確認者", "確認日期": "2026-09-08", "確認狀態": "已確認",
    }.items():
        sheet.cell(2, columns[label], value)
    workbook.save(path)
    workbook.close()


def test_confirmations_persist_and_only_fill_matching_case(tmp_path: Path) -> None:
    original = records()
    path = tmp_path / "confirm.xlsx"
    sync_review(original, path)
    approve(path)
    updated = apply_confirmations(original, path)
    assert updated["final_stage"].tolist() == ["II", ""]
    assert updated.loc[0, "final_stage_source"] == "patient_confirmed"
    assert original["final_stage"].tolist() == ["", ""]
    assert build_queue(updated)["病歷號"].tolist() == ["002"]
    sync_review(updated, path)
    assert apply_confirmations(original, path)["final_stage"].tolist() == ["II", ""]
    assert path.with_name("confirm.backup.xlsx").exists()
    changed = original.copy()
    changed.loc[0, "case_id"] = "C"
    assert apply_confirmations(changed, path).loc[0, "final_stage"] == ""
    changed.loc[0, "case_id"] = "A"
    changed.loc[0, "clinical_tnm"] = "cT3N0M0"
    assert apply_confirmations(changed, path).loc[0, "final_stage"] == ""


def test_existing_stage_is_not_overwritten_and_absent_cases_retained(tmp_path: Path) -> None:
    original = records()
    path = tmp_path / "confirm.xlsx"
    sync_review(original, path)
    approve(path)
    original.loc[0, "final_stage"] = "III"
    updated = apply_confirmations(original, path)
    assert updated.loc[0, "final_stage"] == "III"
    sync_review(updated, path)
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb["個案確認"].iter_rows(values_only=True)
    headers = next(rows)
    data = [dict(zip(headers, row)) for row in rows]
    assert next(row for row in data if row["病歷號"] == "001")["本次狀態"] == "來源期別與確認值不同，請核對"
    wb.close()
    sync_review(original.iloc[[1]], path)
    assert apply_confirmations(records(), path).loc[0, "final_stage"] == "II"


def test_queue_ignores_blank_drug_stages_but_includes_missing_tnm() -> None:
    source = records()
    source.loc[0, "final_stage"] = "II"
    drug = source.iloc[[0]].copy()
    drug["clinical_tnm"] = ""
    drug["final_stage"] = ""
    drug["report_name"] = "用藥明細"
    source.loc[1, "clinical_tnm"] = ""
    source.loc[1, "report_name"] = "用藥明細"
    queue = build_queue(pd.concat([source, drug], ignore_index=True))
    assert queue["病歷號"].tolist() == ["002"]
    assert "缺期別及TNM" in queue.iloc[0]["待確認原因"]


def test_queue_keeps_different_case_of_same_patient() -> None:
    source = records()
    source.loc[0, "final_stage"] = "II"
    source.loc[1, "medical_record_no"] = "001"
    source.loc[1, "patient_key"] = "MRN:001"
    queue = build_queue(source)
    assert queue["個案編號"].tolist() == ["B"]


def test_incomplete_approval_or_changed_identity_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "confirm.xlsx"
    sync_review(records(), path)
    approve(path)
    wb = load_workbook(path)
    sheet = wb["個案確認"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, columns["確認人"]).value = None
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="必須填寫"):
        apply_confirmations(records(), path)
    approve(path)
    wb = load_workbook(path)
    wb["個案確認"].cell(2, columns["病歷號"], "999")
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="識別欄被修改"):
        apply_confirmations(records(), path)


def test_ambiguous_case_cannot_be_approved(tmp_path: Path) -> None:
    source = records()
    source.loc[0, ["case_id", "diagnosis_date"]] = ""
    path = tmp_path / "confirm.xlsx"
    sync_review(source, path)
    approve(path)
    with pytest.raises(ValueError, match="個案識別不足"):
        apply_confirmations(source, path)


def test_cli_round_trip_updates_report_and_preserves_confirmation(tmp_path: Path, monkeypatch) -> None:
    from report_cleaner import cli
    from report_cleaner.config import build_paths
    from report_cleaner.patient_stage_review import review_path

    paths = build_paths(tmp_path)
    shutil.copytree(build_paths().mapping_dir, paths.mapping_dir)
    paths.raw_dir.mkdir()
    pd.DataFrame([
        {"病歷號碼": "001", "病患姓名": "測試甲", "個案編號": "A",
         "癌別": "LC_肺癌", "最初診斷日期": "20250101", "組合": "cTxN0M0", "期別": ""},
        {"病歷號碼": "002", "病患姓名": "測試乙", "個案編號": "B",
         "癌別": "LC_肺癌", "最初診斷日期": "20250101", "組合": "cTxN0M0", "期別": ""},
    ]).to_csv(paths.raw_dir / "D1.2 TNM測試.csv", index=False, encoding="utf-8-sig")
    monkeypatch.setattr(cli, "build_paths", lambda: paths)
    cli.main([])
    approve(review_path(tmp_path))
    for _ in range(2):
        cli.main([])
        report = load_workbook(paths.output_file, read_only=True, data_only=True)
        rows = report["病人總表"].iter_rows(values_only=True)
        headers = next(rows)
        patients = {row["病歷號"]: row for row in (dict(zip(headers, values)) for values in rows)}
        assert patients["001"]["最終期別"] == "II"
        assert patients["001"]["期別資料狀態"] == "含個案人工確認"
        assert patients["002"]["最終期別"] is None
        pending = report["病人期別待確認"]
        assert pending.max_row == 2
        assert pending.cell(2, 2).value == "002"
        report.close()
