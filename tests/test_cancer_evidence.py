from pathlib import Path
import shutil

import pandas as pd
from openpyxl import load_workbook

from report_cleaner.cancer_evidence import (
    annotate_evidence, eligible_rows, mark_registry_source, REGISTRY, PATHOLOGY, PENDING,
)
from report_cleaner.cleaner import (
    _build_standardized_frame, _deduplicate_standardized, _enrich_from_patient_lookup,
    _standardize_order_detail,
)


def pathology_row(mrn="001", **updates):
    row = {"patient_key": f"MRN:{mrn}", "medical_record_no": mrn,
           "order_category": "pathology", "pathology_id": "P01", "pathology_result": "惡性",
           "pathology_level": "最終報告", "source_file": "pathology.xls"}
    row.update(updates)
    return row


def test_evidence_requires_positive_report_not_order_or_diagnosis_hint():
    rows = [pathology_row(), pathology_row("002", pathology_result="未見惡性"),
            pathology_row("003", pathology_level="初步報告"),
            {"patient_key": "MRN:004", "medical_record_no": "004", "surgery": "Y",
             "cancer_type": "BR_乳癌", "icd10_code": "C50", "source_file": "order.xls"}]
    result = annotate_evidence(pd.DataFrame(rows))
    assert result["cancer_evidence"].tolist() == [PATHOLOGY, PENDING, PENDING, PENDING]
    assert eligible_rows(result)["medical_record_no"].tolist() == ["001"]


def test_specimen_conflict_is_not_resolved_by_row_order_or_other_specimens():
    rows = [pathology_row(), pathology_row(pathology_result="良性"),
            pathology_row("002"), pathology_row("002", pathology_id="P02", pathology_result="良性")]
    result = annotate_evidence(pd.DataFrame(rows))
    assert result["cohort_eligible"].tolist() == ["", "", "Y", "Y"]
    assert "不同結果" in result.loc[0, "pathology_review_note"]


def test_registry_evidence_survives_negative_pathology_and_does_not_cross_mrn():
    source = pd.DataFrame([{"case_id": "CASE1", "cancer_type": "BR_乳癌", "medical_record_no": "001",
                            "patient_key": "MRN:001", "source_file": "registry.csv"}])
    registry = mark_registry_source(source, "D1.2 TNM")
    data = pd.concat([registry, pd.DataFrame([pathology_row(pathology_result="良性"), pathology_row("0001")])], ignore_index=True)
    result = annotate_evidence(data)
    assert result["cancer_evidence"].tolist() == [REGISTRY, REGISTRY, PATHOLOGY]
    assert mark_registry_source(source, "乳癌手術清單")["registry_evidence"].eq("").all()


def test_order_report_name_and_rule_never_supply_cancer_type():
    frame = pd.DataFrame({"病歷號": ["001"], "醫令代碼": ["63001B"], "醫令名稱": ["測試術式"]})
    row = _standardize_order_detail(frame, "乳房手術", "surgery.xls", "", "").iloc[0]
    assert row["cancer_type"] == row["diagnosis_text"] == ""


def test_different_mrn_does_not_inherit_case_via_matching_name():
    frame = pd.DataFrame({"medical_record_no": ["002", ""], "case_id": ["", ""], "patient_name": ["測試姓名"] * 2})
    lookup = {"CASE1": {"medical_record_no": "001", "patient_name": "測試姓名", "cancer_type": "BR_乳癌"}}
    enriched = _enrich_from_patient_lookup(frame, lookup)
    assert enriched["case_id"].eq("").all()
    assert enriched["cancer_type"].eq("").all()


def test_pathology_standardizer_preserves_specimen_versions_without_treatment_flags():
    source = pd.DataFrame({"medical_record_no": ["001"] * 3, "patient_name": ["測試個案"] * 3,
                           "病理編號": ["P1", "P2", "P1"], "result": ["惡性", "良性", "良性"],
                           "層級": ["最終報告"] * 3, "報告時間": ["20260901123000"] * 3,
                           "細項名稱": ["Surgical pathology"] * 3, "醫令": ["X"] * 3})
    frame = _build_standardized_frame(source, "A4病理", "path.xls", "", "")
    assert len(_deduplicate_standardized(frame)) == 3
    assert frame["cancer_type"].eq("").all()
    assert frame["diagnosis_date"].eq("").all()
    assert frame["surgery"].eq("").all()
    assert frame.loc[0, "pathology_report_time"] == "2026-09-01 12:30:00"


def test_cli_filters_total_summary_and_trend_but_preserves_all_records(tmp_path: Path, monkeypatch):
    from report_cleaner.config import build_paths
    from report_cleaner import cli

    paths = build_paths(tmp_path)
    shutil.copytree(build_paths().mapping_dir, paths.mapping_dir)
    paths.raw_dir.mkdir()
    pd.DataFrame([{"病歷號": "001", "姓名": "測試甲", "個案編號": "A", "癌別": "BR_乳癌",
                   "組合": "cTxN0M0", "最初診斷日期": "20250101"}]).to_csv(
        paths.raw_dir / "D1.2 TNM.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([{"病歷號": "002", "姓名": "測試乙", "病理編號": "P2", "結果": "惡性",
                   "層級": "完整報告", "細項名稱": "Surgical pathology", "報告時間": "20260101120000"}]).to_csv(
        paths.raw_dir / "A4病理.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([{"病歷號": mrn, "姓名": name, "醫令代碼": "63001B", "醫令名稱": "測試手術",
                   "看診日期": "20260201"} for mrn, name in [("001", "測試甲"), ("003", "測試丙")]]).to_csv(
        paths.raw_dir / "乳房手術.csv", index=False, encoding="utf-8-sig")
    monkeypatch.setattr(cli, "build_paths", lambda: paths)
    cli.main([])
    wb = load_workbook(paths.output_file, data_only=True)
    def records(sheet):
        rows = wb[sheet].iter_rows(values_only=True)
        headers = next(rows)
        return [dict(zip(headers, row)) for row in rows]
    assert {r["病歷號"] for r in records("病人總表")} == {"001", "002"}
    assert [r["病歷號"] for r in records("癌症資格待確認")] == ["003"]
    assert len(records("standardized_data")) == 4
    assert len(records("病理紀錄")) == 1
    assert {r["病歷號"] for r in records("治療藥物明細")} == {"001"}
    assert sum(r["治療事件數"] for r in records("月趨勢")) == 1
    assert not any(r["病歷號"] == "003" for r in records("病人期別待確認"))
    summary = records("統計摘要")
    assert sum(r["個案筆數"] for r in summary if r["分類"] == "期別資料狀態（病人人數）") == 2
    assert sum(r["個案筆數"] for r in summary if r["分類"] == "納入依據（全部病人人數）") == 3
    assert next(r["個案筆數"] for r in summary if r["分類"] == "治療別" and r["項目"] == "手術") == 1
    wb.close()


def test_order_only_batch_has_empty_total_and_visible_pending_list(tmp_path: Path, monkeypatch):
    from report_cleaner.config import build_paths
    from report_cleaner import cli

    paths = build_paths(tmp_path)
    shutil.copytree(build_paths().mapping_dir, paths.mapping_dir)
    paths.raw_dir.mkdir()
    pd.DataFrame([{"病歷號": "001", "姓名": "測試甲", "醫令代碼": "63001B", "醫令名稱": "測試手術"}]).to_csv(
        paths.raw_dir / "乳房手術.csv", index=False, encoding="utf-8-sig")
    monkeypatch.setattr(cli, "build_paths", lambda: paths)
    cli.main([])
    wb = load_workbook(paths.output_file, data_only=True)
    assert wb["病人總表"].max_row == 1
    assert wb["治療藥物明細"].max_row == 1
    assert wb["癌症資格待確認"].max_row == 2
    assert wb["病人期別待確認"].max_row == 1
    assert wb["standardized_data"].max_row == 2
    wb.close()
