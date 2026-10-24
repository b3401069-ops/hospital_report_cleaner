from datetime import date
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
import pytest

from report_cleaner.case_manager import build_report, plan_fields, ReviewRules
from report_cleaner.case_manager_sources import (
    Sources, classify_order, deduplicate_events, load_sources, merge_cases, mrn, day,
)
from report_cleaner.case_manager_excel import (
    ACTUAL, PLANNED, export_report, read_manual, sync_manual,
)


@pytest.fixture
def root(tmp_path):
    shutil.copytree(Path(__file__).parents[1] / "mapping", tmp_path / "mapping")
    return tmp_path


def case(case_id="CASE-A", cancer="BR_乳癌", **extra):
    return dict({"_mrn":"0000001234", "_case":case_id, "_cancer":cancer,
        "_period":date(2026,9,30), "_source":"D1.2_202609.xlsx", "_row":4,
        "個案編號":case_id, "病患姓名":"測試個案", "病歷號":"0000001234", "癌別":cancer,
        "最初診斷日期":"2026-08-01", "收案日期":"2026-08-03", "年齡":"60",
        "收案個管師":"測試個管師", "最近看診日":"2026-08-01"},**extra)


def event(code="63014B", category="手術", scope="BR", **extra):
    return dict({"mrn":"0000001234", "date":date(2026,8,4), "category":category,
        "code":code,"name":"測試醫令","scope":scope,"drug":False,"location":"門診",
        "source":"source.xlsx","source_row":4,"raw_date":"2026-08-04"},**extra)


def data():
    return Sources(cases=[case(),case("CASE-B","LC_肺癌")],events=[event(),event("36015B","放療計畫",""),event("DRUG","化學治療","",drug=True)],as_of=date(2026,9,30))


def test_case_grain_plan_actual_and_scope(root):
    sources=data()
    sources.plans=[case(**{"手術填寫日期":"2026-08-02","手術切除方式":"計畫內容"})]
    report=build_report(sources,root)
    assert len(report.actual)==2
    assert report.planned[0]["手術"]
    assert report.actual[0]["手術"]
    assert not report.actual[1]["手術"]
    assert not report.actual[0]["放射治療"]
    assert report.actual[1]["化學治療"]
    assert report.actual[0]["查詢鍵"]=="0000001234#1"
    assert report.actual[1]["查詢鍵"]=="0000001234#2"
    assert "多癌別" in report.actual[0]["資料提醒"]


def test_plan_alone_is_not_delivered(root):
    sources=Sources(cases=[case()], plans=[case(**{"手術填寫日期":"2026-08-02"})],as_of=date(2026,9,30))
    report=build_report(sources,root)
    assert report.planned[0]["手術"]
    assert not report.actual[0]["手術"]
    assert report.actual[0]["首次治療日"] is None


def test_order_only_patient_not_in_cancer_cohort(root):
    sources=data()
    sources.events.append(event(mrn="0000099999"))
    report=build_report(sources,root)
    assert not any(r["病歷號"]=="0000099999" for r in report.actual)
    assert any(r["病歷號"]=="0000099999" for r in report.pending)


def test_pathology_only_patient_is_not_lost(root):
    sources=data()
    sources.pathology=[{"_mrn":"0000099999","_source":"A4.xlsx","結果":"惡性","層級":"完整報告","病理編號":"SPEC-TEST"}]
    report=build_report(sources,root)
    assert not any(r["病歷號"]=="0000099999" for r in report.actual)
    assert any(r["病歷號"]=="0000099999" and "病理" in r["待確認原因"] for r in report.pending)


def test_overlap_dedup_retains_provenance():
    rows=deduplicate_events([event(),event(source="overlap.xlsx"),event(date=date(2026,8,5))])
    assert len(rows)==2
    assert len(rows[0]["references"])==2


def test_planning_order_overrides_legacy_mapping():
    assert classify_order("36015B","",[{"order_code":"36015B","treatment_type":"radiation_therapy"}])==("放療計畫","")
    assert classify_order("36012B","",[])==("放射治療","")
    assert classify_order("UNKNOWN","切除術",[])==("未分類醫令","")


def test_latest_snapshot_and_case_identity():
    old=case(**{"結案日期":"2026-01-01","期別":"II"})
    old["_period"]=date(2026,8,31)
    new=case(**{"結案日期":"","期別":"III"})
    assert merge_cases([new,old])[0]["期別"]=="III"
    assert merge_cases([new,old])[0]["結案日期"]==""
    other=case(); other["_mrn"]="0000009999"
    assert len(merge_cases([new,other]))==2


def test_clinical_and_pathological_stages_are_both_preserved(root):
    clinical=case(**{"組合":"cT2N0M0","期別":"II"})
    pathology=case(**{"組合":"pT3N0M0","期別":"III"})
    pathology["_row"]=5
    sources=Sources(cases=[clinical,pathology],as_of=date(2026,9,30))
    report=build_report(sources,root)
    assert "II" in report.actual[0]["分期"] and "III" in report.actual[0]["分期"]
    assert "cT2" in report.actual[0]["分期"] and "pT3" in report.actual[0]["分期"]
    assert "不一致" not in report.actual[0]["期別資料狀態"]
    pathology["期別"]=""
    report=build_report(sources,root)
    assert "部分" in report.actual[0]["期別資料狀態"]


def test_conflicting_cancer_is_not_arbitrarily_selected(root):
    report=build_report(Sources(cases=[case(),case(cancer="LC_肺癌")],events=[event()],as_of=date(2026,9,30)),root)
    assert report.actual[0]["癌別"]==""
    assert "不同癌別" in report.actual[0]["資料提醒"]
    assert not report.actual[0]["手術"]


def test_no_dates_no_cross_case_guess(root):
    raw=case(**{"最初診斷日期":"","收案日期":""})
    report=build_report(Sources(cases=[raw],events=[event()],as_of=date(2026,9,30)),root)
    assert not report.actual[0]["手術"]
    assert "缺診斷" in report.actual[0]["資料提醒"]


def test_event_window_cutoff_and_missing_date(root):
    sources=data()
    sources.events=[event(date=date(2020,1,1)),event(date=date(2026,10,1)),event(date=None)]
    report=build_report(sources,root)
    assert not report.actual[0]["手術"]
    assert len(report.events)==3
    assert "日期缺漏" in report.actual[0]["資料提醒"]


def test_mrn_and_dates():
    assert mrn("1234")=="0000001234"
    assert mrn("12345678901")=="12345678901"
    assert day("1151001")==date(2026,10,1)
    assert day("1150230") is None


@pytest.mark.parametrize("name,start,end",[
    ("乳癌手術_11204-11206.xlsx","2023-04-01","2023-06-30"),
    ("化療_11207-11209.xlsx","2023-07-01","2023-09-30"),
    ("CASWQ02_1100101-1150930.xlsx","2021-01-01","2026-09-30"),
    ("D1.2_202101-202609.xlsx","2021-01-01","2026-09-30"),
    ("A1.41_20260901-20260930.xlsx","2026-09-01","2026-09-30"),
])
def test_source_period_does_not_join_roc_digits(name,start,end):
    from report_cleaner.cleaner import _extract_period_from_filename
    assert _extract_period_from_filename(name)==(start,end)


def test_manual_roundtrip_and_new_event_not_resolved(root):
    sources=data()
    report=build_report(sources,root)
    path=root/"review/個管人工確認.xlsx"
    records=sync_manual(report,path)
    assert records
    w=load_workbook(path)
    s=w["人工確認紀錄"]
    key=s["H2"].value
    s["D2"],s["E2"],s["F2"]="已確認無異常","測試確認人",date(2026,10,1)
    w.save(path)
    refreshed=build_report(sources,root)
    sync_manual(refreshed,path)
    assert next(t for t in refreshed.tasks if t["確認鍵"]==key)["確認結果"]=="已確認無異常"
    sources.cases[0]["最初診斷日期"]="2026-09-01"
    changed=build_report(sources,root)
    sync_manual(changed,path)
    assert all(not t["確認結果"] for t in changed.tasks if t["個案編號"]=="CASE-A")
    assert any(r["比對鍵（自動）"]==key for r in read_manual(path))


def test_manual_invalid_completion_fails(root):
    report=build_report(data(),root)
    path=root/"review/manual.xlsx"
    sync_manual(report,path)
    w=load_workbook(path); w["人工確認紀錄"]["D2"]="已確認無異常"; w.save(path)
    with pytest.raises(ValueError,match="確認人"):
        read_manual(path)


def test_case_stage_confirmation_roundtrip(root):
    from report_cleaner.patient_stage_review import sync_review, review_path, INPUTS, COLUMNS
    sources=Sources(cases=[case(**{"組合":"cTxNxMx"})],as_of=date(2026,9,30))
    report=build_report(sources,root)
    path=review_path(root)
    sync_review(report.stages,path)
    w=load_workbook(path); s=w["個案確認"]
    for label,value in zip(INPUTS,["II","臨床","測試依據","測試人員","2026-10-01","已確認"]):
        s.cell(2,COLUMNS.index(label)+1,value)
    w.save(path)
    report=build_report(sources,root)
    assert "期別：II" in report.actual[0]["分期"]
    assert "個案人工確認" in report.actual[0]["期別資料狀態"]
    assert not any(t["待辦類型"]=="期別未填" for t in report.tasks)
    sources.cases[0]["組合"]="cT1NxMx"
    report=build_report(sources,root)
    assert "期別待確認" in report.actual[0]["分期"]


def test_unknown_reported_stage_is_not_silently_inferred(root):
    (root/"mapping/tnm_stage_mapping.csv").write_text(
        "cancer_type,tnm_pattern,final_stage,notes\nBR_乳癌,cT2N0M0,II,test\n",encoding="utf-8-sig")
    sources=Sources(cases=[case(**{"組合":"cT2N0M0","期別":"N"})],as_of=date(2026,9,30))
    report=build_report(sources,root)
    assert "來源期別代碼待核對：N" in report.actual[0]["分期"]
    assert "期別：II" not in report.actual[0]["分期"]


def test_read_legacy_xls_matches_valid_row_count():
    from report_cleaner.case_manager_sources import read_table
    sources=list((Path(__file__).parents[1]/"raw").glob("*.xls"))
    if not sources:
        pytest.skip("No local legacy reports")
    rows,_,_=read_table(sources[0])
    assert rows
    assert all(r["_mrn"].isdigit() for r in rows)


def test_format_and_formula_safety(root):
    sources=data()
    sources.cases[0]["病患姓名"]='=HYPERLINK("https://invalid","x")'
    report=build_report(sources,root)
    manual=root/"review/manual.xlsx"
    records=sync_manual(report,manual)
    output=root/"case_report.xlsx"
    export_report(report,output,records,manual)
    w=load_workbook(output)
    assert [c.value for c in w["病人預定治療總表"][1]]==PLANNED
    assert [c.value for c in w["病人實際治療總表"][1]]==ACTUAL
    assert w["病人實際治療總表"]["G2"].data_type=="s"
    assert w["病人實際治療總表"].freeze_panes=="I2"
    assert "AY$3" in w["病人查詢"]["I3"].value
    assert "INVALID" in w["病人查詢"]["H1"].value
    assert w["病人查詢"]["B6"].data_type=="f"
    assert w.active.title=="個管待辦清單"
    assert w["治療醫令明細"].max_row==4


def test_recursive_sources_skip_merged_outputs(root):
    raw=root/"raw/sub"; raw.mkdir(parents=True)
    w=Workbook(); s=w.active
    s.append(["報表標題"])
    s.append(["病歷號碼","個案編號","癌別","病患姓名","收案日期"])
    s.append(["1234","CASE-A","BR_乳癌","測試個案","1150801"])
    w.save(raw/"D1.2_202609.xlsx")
    w.save(raw/"A1.41_合併_202609.xlsx")
    sources=load_sources(root/"raw",root/"mapping")
    assert len(sources.cases)==1
    assert sources.cases[0]["_mrn"]=="0000001234"
    assert any(r["類型"]=="排除" for r in sources.audit)
