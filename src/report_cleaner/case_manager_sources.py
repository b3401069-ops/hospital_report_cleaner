"""Read Hippo exports without treating a treatment order as a cancer diagnosis."""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

from .cleaner import (
    _classify_drug_treatment, _extract_report_period, _load_drug_treatment_mapping,
    _load_order_treatment_mapping, _normalize_cancer_type_value, _normalize_date_value,
)

TREATMENTS = {
    "surgery": "手術", "radiation_therapy": "放射治療", "chemotherapy": "化學治療",
    "targeted_therapy": "標靶治療", "immunotherapy": "免疫治療",
    "hormone_therapy": "抗荷爾蒙治療", "tace": "肝動脈栓塞（TACE）", "rfa": "射頻燒灼（RFA）",
}
MRN_HEADERS = ("病歷號", "病歷號碼", "病歷編號")
SCOPES = {"乳癌": "BR", "肺癌": "LC", "結腸": "CR", "直腸": "CR", "肝癌": "LI", "攝護腺": "PR", "膀胱": "BL"}


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def mrn(value) -> str:
    value = re.sub(r"\.0$", "", text(value))
    return value.zfill(10) if value.isdigit() and len(value) <= 10 else value


def day(value) -> date | None:
    value = text(value)
    if not value:
        return None
    parsed = pd.to_datetime(_normalize_date_value(value), errors="coerce")
    return None if pd.isna(parsed) else parsed.date()


def first(row, *names):
    return next((row.get(n, "") for n in names if row.get(n, "")), "")


def cancer_scope(value):
    value = text(value)
    for word, code in SCOPES.items():
        if word in value:
            return code
    return value.split("_")[0] if "_" in value else ""


@dataclass
class Sources:
    cases: list[dict] = field(default_factory=list)
    plans: list[dict] = field(default_factory=list)
    events: list[dict] = field(default_factory=list)
    pathology: list[dict] = field(default_factory=list)
    audit: list[dict] = field(default_factory=list)
    as_of: date | None = None


def read_table(path: Path):
    legacy = path.suffix.lower() == ".xls"
    book = pd.ExcelFile(path) if legacy else load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = None if legacy else book.worksheets[0]
        preview = book.parse(header=None, nrows=8).fillna("").values.tolist() if legacy else list(sheet.iter_rows(max_row=8, values_only=True))
        def score(row):
            values = {text(c) for c in row}
            return 10 * len(values.intersection(MRN_HEADERS)) + len(values.intersection(
                {"個案編號", "醫令代碼", "醫令碼", "病患姓名", "姓名", "病理編號", "癌別"}))
        index = max(range(len(preview)), key=lambda i: score(preview[i]), default=0)
        if not preview or score(preview[index]) < 10:
            return [], "", ""
        headers = [text(c) for c in preview[index]]
        rows = []
        values_rows = book.parse(header=None, skiprows=index+1).fillna("").itertuples(index=False,name=None) if legacy else sheet.iter_rows(min_row=index+2,values_only=True)
        for number, values in enumerate(values_rows, index + 2):
            row = {h: text(v) for h, v in zip(headers, values) if h}
            ident = mrn(first(row, *MRN_HEADERS))
            if not ident or not ident.isdigit() or len(ident) > 10:
                continue
            row.update(_mrn=ident, _row=number)
            rows.append(row)
        start, end = _extract_report_period(path.name, pd.DataFrame(preview))
        return rows, start, end
    finally:
        book.close()


def rule_matches(code, name, rule):
    return bool((rule.get("order_code") and code == rule["order_code"].upper())
                or (rule.get("order_code_prefix") and code.startswith(rule["order_code_prefix"].upper()))
                or (rule.get("name_pattern") and re.search(rule["name_pattern"], name, re.I)))


def classify_order(code, name, rules):
    # Planning orders are not delivered fractions, even if the legacy mapping
    # labels them radiation_therapy.
    if code == "36015B":
        return "放療計畫", ""
    if code in {"36012B", "36022B"}:
        return "放射治療", ""
    if code == "50011C":
        return "膀胱灌注", "BL"
    matches = [r for r in rules if rule_matches(code, name, r)]
    kinds = {TREATMENTS[r["treatment_type"]] for r in matches if r["treatment_type"] in TREATMENTS}
    scopes = {cancer_scope(r.get("cancer_type", "")) for r in matches} - {""}
    if len(kinds) != 1 or len(scopes) > 1:
        return "未分類醫令", ""
    kind = next(iter(kinds))
    scope = next(iter(scopes), "")
    if kind in {"肝動脈栓塞（TACE）", "射頻燒灼（RFA）"}:
        scope = "LI"
    return kind, scope


def load_sources(source: Path, mapping: Path) -> Sources:
    if not source.is_dir():
        raise ValueError("找不到來源資料夾。請以 --source 指定 Hippo 或將原始報表放入 raw。")
    result = Sources()
    drug_rules = _load_drug_treatment_mapping(mapping / "drug_treatment_mapping.csv")
    order_rules = _load_order_treatment_mapping(mapping / "order_treatment_mapping.csv")
    # Classify once per code/name, not once per administration.
    drug_cache = {}
    for path in sorted(p for p in source.rglob("*") if p.is_file() and p.suffix.lower() in {".xlsx", ".xls", ".xlsm"}):
        relative = path.relative_to(source).as_posix()
        name = path.name
        reason = ""
        if name.startswith("~$") or any(x in name for x in ["比較", "_vs_", "合併", "未完成", "cleaned", "整合查詢", "病人期別確認", "個管人工確認"]):
            reason = "排除暫存／比較／合併成果（使用原始分檔）"
        if reason:
            result.audit.append({"來源檔案": relative, "類型": "排除", "有效列數": 0, "說明": reason})
            continue
        rows, start, end = read_table(path)
        if not rows:
            result.audit.append({"來源檔案": relative, "類型": "無可用資料", "有效列數": 0, "說明": "未辨識病歷號表頭或沒有有效病歷號"})
            continue
        fields = set(rows[0])
        if "病理編號" in fields:
            kind = "病理"
        elif re.search(r"D1\.4[1-8]", name):
            kind = "計畫"
        elif "個案編號" in fields and ("D1.2" in name or "個案清單" in name or "CASWQ" in name):
            kind = "收案"
        elif "醫令代碼" in fields:
            kind = "醫令"
        elif "醫令碼" in fields:
            kind = "藥物"
        else:
            kind = "未支援"
        result.audit.append({"來源檔案": relative, "類型": kind, "有效列數": len(rows), "說明": "", "期間起": start, "期間迄": end})
        if kind == "未支援":
            continue
        period = day(end) or date.min
        if period != date.min:
            result.as_of = max(result.as_of or period, period)
        for row in rows:
            row.update(_source=relative, _period=period, _kind=kind)
            if kind in {"收案", "計畫"}:
                row["_case"] = first(row, "個案編號")
                row["_cancer"] = _normalize_cancer_type_value(first(row, "癌別", "癌別名稱"))
                (result.cases if kind == "收案" else result.plans).append(row)
            elif kind == "病理":
                result.pathology.append(row)
            else:
                code = first(row, "醫令代碼", "醫令碼").upper()
                title = first(row, "醫令名稱", "名稱")
                if not code:
                    continue
                if kind == "藥物":
                    if (code, title) not in drug_cache:
                        flags = _classify_drug_treatment(pd.Series([code]), pd.Series([title]), drug_rules)
                        types = [TREATMENTS[k] for k, v in flags.items() if v.iloc[0] == "Y"]
                        drug_cache[code, title] = types[0] if len(types) == 1 else "未分類藥物"
                    category, scope = drug_cache[code, title], ""
                else:
                    category, scope = classify_order(code, title, order_rules)
                when = day(first(row, "執行日期-起", "開單日期", "看診日期"))
                result.events.append({"mrn": row["_mrn"], "date": when, "category": category,
                    "code": code, "name": title, "scope": scope, "drug": kind == "藥物",
                    "location": first(row, "來源", "類別"), "source": relative,
                    "source_row": row["_row"], "raw_date": first(row, "執行日期-起", "開單日期", "看診日期")})
    if not result.cases:
        raise ValueError("來源中找不到 D1.2／CASWQ02／個案清單。不能僅用手術或藥物名單建立癌症個案。")
    return result


def deduplicate_events(events):
    """One MRN/date/code/location per event summary; retain all source references.

    These exports have no stable administration ID: this is a presence/date
    summary, never a dose, billing or fraction-count ledger.
    """
    unique = {}
    for e in events:
        key = (e["mrn"], e["date"], e["code"], e["location"], e["drug"])
        if e["date"] is None:
            key += (e["source"], e["source_row"])
        if key not in unique:
            unique[key] = dict(e, references=[])
        target = unique[key]
        reference = f'{e["source"]}（列{e["source_row"]}）'
        if reference not in target["references"]:
            target["references"].append(reference)
    return sorted(unique.values(), key=lambda e: (e["mrn"], e["date"] or date.min, e["code"]))


def merge_cases(rows):
    groups = defaultdict(list)
    for row in rows:
        if row["_case"]:
            groups[row["_mrn"], row["_case"]].append(row)
    merged = []
    for key, group in sorted(groups.items()):
        group.sort(key=lambda r: (r["_period"], "D1.2" in r["_source"], r["_source"], r["_row"]))
        record = dict(group[-1])
        # Complete missing fields from previous exports, but do not resurrect
        # a closure/death that the newest snapshot has explicitly cleared.
        for old in reversed(group[:-1]):
            for field_name, value in old.items():
                if field_name not in {"結案日期", "結案原因", "死亡日期", "期別", "組合"} and not record.get(field_name):
                    record[field_name] = value
        record["_sources"] = sorted({r["_source"] for r in group})
        latest_group = [r for r in group if (r["_period"],r["_source"]) == (group[-1]["_period"],group[-1]["_source"])]
        cancers = {r["_cancer"] for r in latest_group if r["_cancer"]}
        record["_cancer_conflict"] = len(cancers) > 1
        if record["_cancer_conflict"]:
            record["_cancer"] = ""
            record["癌別"] = ""
        tnm_rows = [r for r in group if "D1.2" in r["_source"]]
        if tnm_rows:
            latest = max((r["_period"], r["_source"]) for r in tnm_rows)
            tnm_rows = [r for r in tnm_rows if (r["_period"], r["_source"]) == latest]
            informative = [r for r in tnm_rows if r.get("組合") or r.get("期別")]
            tnm_rows = informative or tnm_rows[:1]
        record["_tnm_rows"] = tnm_rows or [dict(record)]
        values = defaultdict(set)
        for row in tnm_rows:
            if row.get("期別"):
                values[row.get("組合", "")].add(row["期別"])
        record["_stage_conflict"] = any(len(stages) > 1 for stages in values.values())
        merged.append(record)
    return merged
