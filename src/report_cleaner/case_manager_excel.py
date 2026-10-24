"""The case managers' two-table workbook layout, generated without a PHI template."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
import os
import shutil
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

from .case_manager import CaseReport
from .case_manager_sources import day, text

PLANNED = "收案日期|病患姓名|病歷號|性別|年齡|癌別|分期|期別資料狀態|納入依據|最後回診日|最後資料日期|主責醫師|個管師|手術|放射治療|化學治療|標靶治療|免疫治療|抗荷爾蒙治療|肝動脈栓塞（TACE）|射頻燒灼（RFA）|藥物種類數|資料提醒|個案編號|診斷日|診斷名稱|CCRT|預定緩和治療|治療摘要|藥物摘要|計畫書日期|計畫資料來源|癌症個案數|查詢鍵".split("|")
ACTUAL = "收案日期|個案分類|結案日期/原因|死亡日期|必要提報事件|是否符合計畫書|病患姓名|病歷號|性別|年齡|癌別|分期|期別資料狀態|納入依據|最後回診日|最後資料日期|最近追蹤日|追蹤狀態|失聯狀態|主責醫師|個管師|手術|放射治療|抗荷爾蒙治療|化學治療|標靶治療|免疫治療|肝動脈栓塞（TACE）|射頻燒灼（RFA）|藥物種類數|資料提醒|首次治療日|首次治療項目|診斷至首次治療天數|首次手術日|術後至化療天數|術後至放療天數|放療完成度|放療中斷|MDT討論|計畫書後新療程|個案編號|診斷日|診斷名稱|CCRT|治療摘要|藥物摘要|癌症個案數|收案年度|收案季|查詢鍵".split("|")
TASKS = "嚴重度|個管師|病歷號|個案編號|癌別|待辦類型|說明|相關日期|逾期天數|建議動作|確認結果|確認備註".split("|")
MANUAL = "個案編號|病歷號|待辦類型|確認結果|確認人|確認日期|備註|比對鍵（自動）".split("|")
MANUAL_CONTEXT = ["相關日期", "待辦說明", "本次狀態"]
OPTIONS = ["已確認無異常", "外院執行", "病人拒絕治療", "計畫變更", "已補登資料", "已聯繫追蹤", "已安排MDT", "確實異常-處理中", "確實異常-已提報", "其他（見備註）"]


def atomic_save(book, path, *, backup=False):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(suffix=".xlsx", prefix=".writing-", dir=path.parent)
    os.close(fd)
    try:
        book.save(temp)
        if backup and path.exists():
            folder = path.parent / "backups"
            folder.mkdir(exist_ok=True)
            shutil.copy2(path, folder / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx")
        os.replace(temp, path)
    finally:
        Path(temp).unlink(missing_ok=True)


def set_value(cell, value):
    cell.value = value
    # External text must never become an executable Excel formula.
    if isinstance(value, str):
        cell.data_type = "s"
    if isinstance(value, (datetime, date)):
        cell.number_format = "yyyy-mm-dd"


def table(book, title, headers, records, freeze="A2"):
    sheet = book.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = freeze
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(1, col, header)
        color = "526677"
        if header in {"分期", "期別資料狀態", "癌別", "納入依據"}:
            color = "376E8A"
        elif header in {"手術", "放射治療", "化學治療", "標靶治療", "免疫治療", "抗荷爾蒙治療", "肝動脈栓塞（TACE）", "射頻燒灼（RFA）"}:
            color = "447765"
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name="Microsoft JhengHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = 14
        if any(t in header for t in ["摘要", "提醒", "說明", "原因", "來源", "依據", "計畫", "狀態", "完成", "中斷"]) or header in {"必要提報事件", "建議動作", "名稱"}:
            width = 36
        elif header in {"手術", "放射治療", "化學治療", "標靶治療", "免疫治療", "抗荷爾蒙治療"}:
            width = 36
        elif header in {"確認結果", "待辦類型", "分期", "MDT討論"}:
            width = 26
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[1].height = 42
    body_font = Font(name="Microsoft JhengHei", size=10, color="243441")
    text_alignment = Alignment(vertical="top",wrap_text=True,indent=1,horizontal="left")
    number_alignment = Alignment(vertical="top",wrap_text=True,indent=1,horizontal="right")
    stripe = PatternFill("solid",fgColor="F0F4F6")
    for index, record in enumerate(records, 2):
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(index, col)
            set_value(cell, record.get(header))
            cell.font = body_font
            cell.alignment = text_alignment if isinstance(cell.value,str) else number_alignment
            if index % 2 == 0:
                cell.fill = stripe
            if header in {"病歷號", "個案編號"}:
                cell.number_format = "@"
        sheet.row_dimensions[index].height = 46
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2,len(records)+1)}"
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    return sheet


def read_manual(path):
    if not Path(path).exists():
        return []
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if "人工確認紀錄" not in book.sheetnames:
            raise ValueError("人工確認檔缺少「人工確認紀錄」分頁。")
        rows = book["人工確認紀錄"].iter_rows(values_only=True)
        headers = next(rows, ())
        if list(headers[:8]) != MANUAL:
            raise ValueError("人工確認檔前8欄格式已改動，請保留原欄位。")
        records, seen = [], set()
        for values in rows:
            record = {h:text(v) for h,v in zip(headers,values) if h in MANUAL+MANUAL_CONTEXT}
            if not any(record.get(h) for h in MANUAL):
                continue
            key = record.get("比對鍵（自動）", "")
            if not key or len(key) != 64 or key in seen:
                raise ValueError("人工確認檔比對鍵遺失、重複或不是新版確認檔，請勿自行複製識別欄。")
            seen.add(key)
            if record.get("確認結果"):
                if record["確認結果"] not in OPTIONS or not record.get("確認人") or not day(record.get("確認日期")):
                    raise ValueError("請填有效確認結果、確認人與確認日期，再重新執行。")
                if any(v.startswith("=") for v in record.values()):
                    raise ValueError("人工確認請填文字／日期，不可使用公式。")
            records.append(record)
        return records
    finally:
        book.close()


def manual_sheet(book, records, *, editable):
    sheet = table(book, "人工確認紀錄", MANUAL+MANUAL_CONTEXT, records)
    for index in range(2, len(records)+2):
        for col in range(4, 8):
            sheet.cell(index, col).fill = PatternFill("solid", fgColor="FFF2CC" if editable else "EAEFF2")
    sheet.column_dimensions["H"].hidden = True
    sheet.column_dimensions["J"].width = 60
    sheet.column_dimensions["N"].width = 25
    sheet.column_dimensions["Q"].width = 60
    sheet["N1"] = "確認結果選項"
    for i, value in enumerate(OPTIONS, 2):
        sheet.cell(i, 14, value)
    sheet["Q1"] = "只填黃色D:G欄，先看I:K事件日期、說明及本次狀態，存檔關閉後重跑。" if editable else "這是確認紀錄快照；請到 review/個管人工確認.xlsx 填寫。"
    validation = DataValidation(type="list", formula1="=$N$2:$N$11", allow_blank=True)
    validation.errorTitle = "請選擇確認結果"
    validation.error = "請使用清單選項，細節填備註。"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"D2:D{max(2,len(records)+1)}")
    return sheet


def sync_manual(report, path):
    active = {t["確認鍵"] for t in report.tasks}
    records = {r["比對鍵（自動）"]: r for r in read_manual(path)
               if r["比對鍵（自動）"] in active or any(r.get(h) for h in MANUAL[3:7])}
    for record in records.values():
        record["本次狀態"] = "非本批待辦，保留過去確認"
    for task in report.tasks:
        key = task["確認鍵"]
        identity = {h: task[h] for h in ["個案編號", "病歷號", "待辦類型"]}
        if key in records and any(records[key].get(h) != value for h, value in identity.items()):
            raise ValueError("人工確認檔的病歷號／個案編號／待辦類型被改動，請只填D:G確認欄。")
        record = records.setdefault(key, dict(identity, **{"比對鍵（自動）": key}))
        record.update({"相關日期":task["相關日期"], "待辦說明":task["說明"], "本次狀態":"本批待辦"})
        task["確認結果"] = record.get("確認結果", "")
        task["確認備註"] = record.get("備註", "")
    book = Workbook()
    book.remove(book.active)
    manual_sheet(book, list(records.values()), editable=True)
    atomic_save(book, path, backup=True)
    return list(records.values())


QUERY_ROWS = {
    6:"病患姓名", 7:"性別", 8:"年齡", 9:"癌別", 10:"個案編號", 12:"診斷日", 13:"診斷名稱",
    14:"期別資料狀態", 15:"分期", 16:"納入依據", 17:"手術", 18:"放射治療", 19:"化學治療",
    20:"標靶治療", 21:"免疫治療", 22:"抗荷爾蒙治療", 23:"肝動脈栓塞（TACE）", 24:"射頻燒灼（RFA）",
    25:"CCRT", 26:"治療摘要", 27:"藥物種類數", 28:"藥物摘要", 30:"最後回診日", 31:"最後資料日期",
    32:"最近追蹤日", 33:"追蹤狀態", 34:"失聯狀態", 35:"首次治療項目", 36:"首次治療日",
    37:"診斷至首次治療天數", 38:"術後至化療天數", 39:"術後至放療天數", 40:"放療完成度",
    41:"放療中斷", 42:"MDT討論", 43:"計畫書後新療程", 44:"主責醫師", 45:"個管師",
    46:"癌症個案數", 47:"個案分類", 48:"結案日期/原因", 49:"死亡日期", 50:"必要提報事件",
    51:"是否符合計畫書", 52:"計劃書是否符合院內指引", 53:"資料提醒",
}


def query_sheet(book, count):
    sheet = book.create_sheet("病人查詢")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B5"
    sheet.column_dimensions["A"].width = 27
    for c in "BC":
        sheet.column_dimensions[c].width = 58
    sheet.column_dimensions["D"].width = 34
    sheet.column_dimensions["E"].width = 12
    for c in "HI":
        sheet.column_dimensions[c].hidden = True
    sheet.append(["病人查詢", "預定治療計畫書", "實際治療情形", "資料來源"])
    sheet["A2"] = "輸入病歷號；同人多癌別請切換個案序號。長內容可點儲存格查看完整文字。"
    sheet.merge_cells("A2:E2")
    sheet["A3"], sheet["D3"], sheet["E3"] = "病歷號", "個案序號", 1
    sheet["B3"].number_format = "@"
    sheet["B3"].fill = sheet["E3"].fill = PatternFill("solid", fgColor="FFF2CC")
    digits_removed = "TRIM(B3)"
    for digit in "0123456789":
        digits_removed = f'SUBSTITUTE({digits_removed},"{digit}","")'
    sheet["H4"] = "=" + digits_removed
    sheet["H1"] = '=IF(TRIM(B3)="","",IF(AND(LEN(TRIM(B3))<=10,H4=""),RIGHT("0000000000"&TRIM(B3),10),"INVALID"))'
    end = max(2, count+1)
    sheet["H2"] = f'=IF(OR(H1="",H1="INVALID"),0,COUNTIF(\'病人預定治療總表\'!$C$2:$C${end},H1))'
    for helper, title, keycol in [("H3", "病人預定治療總表", "AH"), ("I3", "病人實際治療總表", "AY")]:
        sheet[helper] = f'=IF($H$2=0,0,IFERROR(MATCH($H$1&"#"&IF($E$3="",1,$E$3),\'{title}\'!${keycol}$2:${keycol}${end},0),0))'
    sheet["B4"] = '=IF(H1="","請輸入病歷號",IF(H1="INVALID","病歷號格式不正確",IF(H2=0,"查無收案個案，請查癌症資格待確認",IF(H3=0,"此個案序號不存在","共"&H2&"個癌症個案，目前第"&IF(E3="",1,E3)&"個"))))'
    sheet["C4"] = "=B4"
    for row, label in QUERY_ROWS.items():
        sheet.cell(row, 1, label)
        origin = "D1.2／個案清單"
        if row in range(17, 29) or row in range(35, 44) or row in {50,51,52}:
            origin = "計畫書／醫令明細；判讀限制見使用說明"
        elif row in {14,15}:
            origin = "D1.2／期別對照／個案確認"
        sheet.cell(row, 4, origin)
        for col, headers, title, helper in [(2, PLANNED, "病人預定治療總表", "$H$3"), (3, ACTUAL, "病人實際治療總表", "$I$3")]:
            if label not in headers:
                if row == 52:
                    sheet.cell(row, col, "需由醫療端核對院內指引，未自動判定")
                continue
            letter = get_column_letter(headers.index(label)+1)
            reference = f"INDEX('{title}'!${letter}$2:${letter}${end},{helper})"
            sheet.cell(row, col, f'=IF({helper}=0,"",IF({reference}="","",{reference}))')
            if label in {"診斷日", "最後回診日", "最後資料日期", "最近追蹤日", "首次治療日", "死亡日期"}:
                sheet.cell(row, col).number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row].height = 60 if row in [17,18,19,20,21,22,26,28,51,53] else 30
    for row in sheet.iter_rows(min_row=1, max_row=53, max_col=5):
        for cell in row:
            cell.font = Font(name="Microsoft JhengHei", size=10, color="243441")
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)
            if cell.row in QUERY_ROWS and cell.column <= 4:
                cell.border = Border(bottom=Side(style="hair",color="D9E1E8"))
                if cell.column == 1:
                    cell.fill = PatternFill("solid",fgColor="EAF0F4")
                    cell.font = Font(name="Microsoft JhengHei",size=10,bold=True,color="243441")
    for col, color in [(2,"E6EEF3"),(3,"E3ECE8")]:
        sheet.cell(1,col).fill = PatternFill("solid",fgColor=color)
        sheet.cell(1,col).font = Font(name="Microsoft JhengHei",size=13,bold=True)
    sheet["A1"].font = Font(name="Microsoft JhengHei",size=16,bold=True)
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[4].height = 38
    validation = DataValidation(type="whole",operator="between",formula1=1,formula2=999,allow_blank=True)
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add("E3")


def export_report(report: CaseReport, output: Path, manual_records, manual_path):
    book = Workbook()
    book.remove(book.active)
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    instructions = [
        "個管版：同一病人可有多個癌症個案；一個個案一列，預定／實際兩表同序。",
        f"統計截止日：{report.as_of}。提醒以來源涵蓋日為基準，不以開檔當天推算。",
        "病人查詢：B3輸入病歷號（可省略前導零）；E3切換同病人的個案序號。",
        "治療格顯示本批醫令日期範圍及名稱；不是完整劑量／分次帳。長內容請展開列高或看明細。",
        "預定治療只來自計畫書；填寫日期不代表實際治療日期。實際表的有醫令也不保證已施作。",
        "36015B只列放療計畫，不計為實際放療；未分類藥物／醫令仍保留明細。",
        "沒有查到醫令不代表沒有治療，可能在外院、來源不足或藥物尚未分類。",
        "以病歷號＋個案編號合併收案；只憑姓名、手術檔名不納入癌症。",
        "治療連結基準：診斷日（缺時用收案日）前90天起；特定部位醫令另限制癌別。",
        "同人多癌別的非部位醫令可能出現在多個個案，不能跨個案加總為醫令人次。",
        "重疊匯出依病歷號、日期、醫令碼、門住別去重；同日同碼多次不能據此計算劑量或放療分次。",
        "明細來源欄保留原始檔名及列號。合併／比較／未完成成果不作為來源，請保留各原始分檔。",
        "癌別未填者保留總表並標示資格待確認，但不納入癌別統計；醫令未連結收案者另列資格待確認。",
        "有病理惡性紀錄不等於已確認是哪個癌症個案；沒有該紀錄也不等於非癌症。",
        "期別確認請填 review/病人期別確認.xlsx；只補同個案、同TNM的空白期別，不覆蓋既有期別。",
        f"待辦確認請填 {manual_path.name}（專案review資料夾），只填黃色D:G，存檔關閉後重跑。",
        "本報表人工確認紀錄是快照，不是永久輸入位置。確認待辦不會自動改期別或醫療紀錄。",
        "舊確認只沿用同一事件／條件的確認鍵；新資料改變條件會重新待確認。已確認待辦保留供稽核。",
        "提醒閾值依提供的範本：首療30天、術後化療56天、未回診90天、追蹤6個月；僅供核對。",
        "放療完成度／中斷、必要提報、MDT適用性、新療程與CCRT尚需醫療端確認，未輸出自動正常結論。",
        "本版未提供國定假日自動判斷；不得據空白假日表計算放療中斷工作日。",
        "統計人數與個案數不同；癌別年齡分布以收案個案計，不是獨立病人人數。",
        "重新清洗必須提供完整原始報表；確認檔不保存五年來源資料。",
    ]
    guide = table(book,"使用說明",["操作與資料限制"],[{"操作與資料限制": x} for x in instructions])
    guide.column_dimensions["A"].width = 125
    todo = table(book,"個管待辦清單",TASKS,report.tasks,"D2")
    summary = table(book,"統計摘要",["項目","數值","說明"],[])
    query_sheet(book,len(report.actual))
    for title, headers, records, freeze in [("病人預定治療總表",PLANNED,report.planned,"D2"),("病人實際治療總表",ACTUAL,report.actual,"I2")]:
        sheet = table(book,title,headers,records,freeze)
        sheet.column_dimensions[get_column_letter(len(headers))].hidden = True
        sheet.sheet_properties.tabColor = "447765" if "實際" in title else "376E8A"
    detail = [{"病歷號":e["mrn"],"日期":e["date"],"治療類別":e["category"],"醫令碼":e["code"],"名稱":e["name"],"門住別":e["location"],"資料來源":join_refs(e["references"]),"只計入癌別":e["scope"] or "無部位限制（多癌別需核對）"} for e in report.events]
    table(book,"治療醫令明細","病歷號|日期|治療類別|醫令碼|名稱|門住別|資料來源|只計入癌別".split("|"),detail,"B2")
    valid = [r for r in report.actual if r["癌別"]]
    counts = Counter((r["收案年度"],r["癌別"]) for r in valid)
    year_totals = Counter(r["收案年度"] for r in valid)
    stats=[]
    for rank, ((year, cancer), number) in enumerate(sorted(counts.items(),key=lambda item:(item[0][0] or 0,-item[1],item[0][1])),1):
        stats.append({"排名":rank,"年度":year,"癌別":cancer,"收案人數":number,"比例":number/year_totals[year]})
    statistic = table(book,"統計資料",["排名","年度","癌別","收案人數","比例"],stats)
    statistic["G1"] = "沿用範本欄名「收案人數」；此處實際口徑是個案數。"
    for i in range(2,len(stats)+2):
        statistic.cell(i,5).number_format="0.0%"
    bins=["0–19歲","20–39歲","40–49歲","50–59歲","60–69歲","70–79歲","80歲以上","未填／不明"]
    ages=defaultdict(Counter)
    for row in valid:
        try:
            age=float(row["年齡"])
            category=next((label for limit,label in zip([20,40,50,60,70,80,131],bins) if 0<=age<limit),bins[-1])
        except (ValueError,TypeError):
            category=bins[-1]
        ages[row["癌別"]][category]+=1
    table(book,"癌別年齡分布",["癌別",*bins],[dict({"癌別":k},**{b:v[b] for b in bins}) for k,v in sorted(ages.items())],"B2")
    manual_sheet(book,manual_records,editable=False)
    holidays=table(book,"國定假日",["日期","名稱","說明"],[])
    holidays["E1"]="未啟用工作日判斷；沒有把週末以外日期視為全部上班日。"
    table(book,"癌症資格待確認",["病歷號","個案編號","待確認原因","來源檔案"],report.pending,"B2")
    table(book,"來源檢查",["來源檔案","類型","有效列數","說明","期間起","期間迄"],report.audit)
    metrics=[("統計截止日",report.as_of,"來源涵蓋日或 --as-of，不以 TODAY() 改變"),
        ("收案個案數",len(report.actual),"同病人可多癌症個案，含資格待確認"),
        ("收案病人人數",len({r["病歷號"] for r in report.actual}),"病歷號去重"),
        ("有癌別個案數",len(valid),"癌別統計分母"),
        ("癌別未填個案數",len(report.actual)-len(valid),"不納入癌別統計"),
        ("醫令去重列數",len(report.events),"非劑量／分次／獨立施作次數"),
        ("待辦項目數",len(report.tasks),"含已確認項目，請篩選確認結果"),
        ("已填確認結果",sum(bool(r["確認結果"]) for r in report.tasks),"處理中不代表問題解除"),
        ("未連結／測試病歷待核對",len(report.pending),"不因有手術醫令而認定癌症")]
    for i,(label,value,note) in enumerate(metrics,2):
        for j,v in enumerate([label,value,note],1):
            cell=summary.cell(i,j)
            set_value(cell,v)
            cell.font=Font(name="Microsoft JhengHei",size=10,color="243441")
            cell.alignment=Alignment(vertical="center",horizontal="right" if j==2 else "left",indent=1,wrap_text=True)
            if i%2==0:
                cell.fill=PatternFill("solid",fgColor="F0F4F6")
            if isinstance(v,int):
                cell.number_format="#,##0"
        summary.row_dimensions[i].height=32
    summary.column_dimensions["A"].width=30
    summary.column_dimensions["C"].width=66
    summary.auto_filter.ref=f"A1:C{len(metrics)+1}"
    count_end=max(2,len(report.tasks)+1)
    todo.conditional_formatting.add(f"K2:K{count_end}",FormulaRule(formula=['$K2=""'],fill=PatternFill("solid",fgColor="FFF2CC")))
    book.active=1
    atomic_save(book,output)
    return Path(output)


def join_refs(values):
    # Excel cell content has a 32767-character hard limit.
    value="\n".join(values)
    if len(value)>32767:
        raise ValueError("單筆來源過多超過Excel儲存格限制，請縮小來源範圍。")
    return value
