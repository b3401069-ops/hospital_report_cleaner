"""Excel live-follow-up formulas and bounded case-level summary tables."""
from datetime import date
from dateutil.relativedelta import relativedelta
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter


def live_followup(sheet, count):
    for r in range(2,count+2):
        base = f'IF(AQ{r}="",A{r},AQ{r})'
        old = f'AND(Q{r}<>"",Q{r}>={base})'
        due = f'IF({old},EDATE(Q{r},6),{base}+30)'
        phase = f'IF({old},"（舊案追蹤）","（新確診）")'
        guard = f'IF(D{r}<>"","免追蹤：已死亡",IF(C{r}<>"","免追蹤：已結案",IF(K{r}="","癌別未填（資格待確認），不判斷",'
        sheet[f'R{r}'] = '='+guard+f'IF(AND(AQ{r}="",A{r}="",Q{r}=""),"缺追蹤基準日",IF(TODAY()>{due},"需追蹤"&{phase}&"：已逾期"&(TODAY()-{due})&"天","未到期"&{phase}&"：請於"&TEXT({due},"yyyy/mm/dd")&"前追蹤")))))'
        sheet[f'S{r}'] = '='+guard+f'IF(O{r}="","無回診紀錄",IF(TODAY()-O{r}>90,"疑似失聯："&(TODAY()-O{r})&"天未回診","未觸發逾90天提醒")))))'
    for col,prefix in [("R","需追蹤"),("S","疑似失聯")]:
        sheet.conditional_formatting.add(f'{col}2:{col}{max(2,count+1)}',FormulaRule(
            formula=[f'LEFT({col}2,{len(prefix)})="{prefix}"'],
            fill=PatternFill('solid',fgColor='FCE4D6'),font=Font(color='9C0006')))
    sheet['F1'].comment = Comment('僅比對計畫書與本批醫令的治療別有無；非指引符合性判定。查無醫令不代表未治療。','報表清洗器')
    sheet['R1'].comment = sheet['S1'].comment = Comment('本欄以TODAY()開檔更新；待辦清單仍是統計截止日快照，不會因開檔自動增加待辦。','報表清洗器')


def summary_tables(sheet, report):
    end=max(2,len(report.actual)+1)
    tasks_end=max(2,len(report.tasks)+1)
    def actual(col): return f"'病人實際治療總表'!${col}$2:${col}${end}"
    def tasks(col): return f"'個管待辦清單'!${col}$2:${col}${tasks_end}"
    def count(criteria):
        return 'COUNTIFS('+','.join(f'{rng},{criterion}' for rng,criterion in criteria)+')'
    row=14
    for mode, values in [
        ('個管師',sorted({r['個管師'] for r in report.actual if r['癌別']})),
        ('癌別',sorted({r['癌別'] for r in report.actual if r['癌別']})),
        ('收案季',[date(report.as_of.year,((report.as_of.month-1)//3)*3+1,1)-relativedelta(months=3*i) for i in reversed(range(8))]),
    ]:
        headers=[mode,'收案個案數','符合計畫書','不符','符合率']
        headers += (['需追蹤','疑似失聯','疑似必要提報','待辦紅件數','待辦黃件數','紅色未確認件數'] if mode=='個管師' else
                    ['第三四期個案','其中經MDT','MDT討論率','有首療天數','首療≤30天','首療≤30天比率','放療中斷個案數','疑似必要提報'] if mode=='癌別' else
                    ['有首療天數','首療≤30天','首療≤30天比率','疑似必要提報'])
        for c,h in enumerate(headers,1):
            cell=sheet.cell(row,c,h); cell.fill=PatternFill('solid',fgColor='376E8A')
            cell.font=Font(name='Microsoft JhengHei',bold=True,color='FFFFFF')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
            sheet.column_dimensions[get_column_letter(c)].width=22 if c==1 else 16
        sheet.row_dimensions[row].height=36
        first_row=row+1
        for value in values:
            row+=1
            label=sheet.cell(row,1,value if mode!='收案季' else f'{value.year-1911}年第{(value.month-1)//3+1}季')
            label.data_type='s'
            base=[(actual('K'),'"<>"')]
            if mode=='收案季':
                stop=value+relativedelta(months=3)
                base += [(actual('A'),f'">="&DATE({value.year},{value.month},1)'),(actual('A'),f'"<"&DATE({stop.year},{stop.month},1)')]
            else:
                base += [(actual('U' if mode=='個管師' else 'K'),f'$A{row}')]
            formulas=[count(base),count(base+[(actual('F'),'"是"')]),count(base+[(actual('F'),'"否*"')]),f'IF(C{row}+D{row}=0,"",C{row}/(C{row}+D{row}))']
            eligible=[(actual('AH'),'">=-30"'),(actual('AH'),'"<=365"'),(actual('AH'),'"<>"')]
            fast=eligible+[(actual('AH'),'"<=30"')]
            reportable=count(base+[(actual('E'),'"疑似必要提報*"')])
            if mode=='個管師':
                task_base=[(tasks('B'),f'$A{row}'),(tasks('E'),'"<>"')]
                formulas += [count(base+[(actual('R'),'"需追蹤*"')]),count(base+[(actual('S'),'"疑似失聯*"')]),reportable,
                    count(task_base+[(tasks('A'),'"紅"')]),count(task_base+[(tasks('A'),'"黃"')]),count(task_base+[(tasks('A'),'"紅"'),(tasks('K'),'""')])]
            elif mode=='癌別':
                # Stage contexts are assessed in Python; MDT status preserves the
                # explicit advanced-stage indication, including recorded meetings.
                advanced=[(actual('M'),'"*第三四期*"')]
                formulas += [count(base+advanced),count(base+advanced+[(actual('AN'),'"有*"')]),f'IF(F{row}=0,"",G{row}/F{row})',
                    count(base+eligible),count(base+fast),f'IF(I{row}=0,"",J{row}/I{row})',
                    count(base+[(actual('AM'),'"*放療中斷（初篩）*"')]),reportable]
            else:
                formulas += [count(base+eligible),count(base+fast),f'IF(F{row}=0,"",G{row}/F{row})',reportable]
            for col,formula in enumerate(formulas,2):
                sheet.cell(row,col,'='+formula)
                if '率' in headers[col-1]: sheet.cell(row,col).number_format='0.0%'
        row+=1
        sheet.cell(row,1,'合計')
        for col in range(2,len(headers)+1):
            letter=get_column_letter(col)
            formula=f'SUM({letter}{first_row}:{letter}{row-1})' if values else '0'
            if col==5: formula=f'IF(C{row}+D{row}=0,"",C{row}/(C{row}+D{row}))'
            elif mode=='癌別' and col==8: formula=f'IF(F{row}=0,"",G{row}/F{row})'
            elif mode=='癌別' and col==11: formula=f'IF(I{row}=0,"",J{row}/I{row})'
            elif mode=='收案季' and col==8: formula=f'IF(F{row}=0,"",G{row}/F{row})'
            sheet.cell(row,col,'='+formula)
            if '率' in headers[col-1]: sheet.cell(row,col).number_format='0.0%'
        row+=3
    for cells in sheet.iter_rows(min_row=14,max_row=row,max_col=13):
        for cell in cells:
            if cell.fill.fgColor.rgb != '00376E8A':
                cell.font=Font(name='Microsoft JhengHei',size=10,color='243441')
                cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[cells[0].row].height=34
