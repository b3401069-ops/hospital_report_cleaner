from __future__ import annotations

from datetime import datetime
from typing import Mapping

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

SHEET_GUIDE = "使用說明"
SHEET_PATIENT_LIST = "病人總表"
SHEET_TREATMENT_DETAIL = "治療藥物明細"
SHEET_QA_SAMPLE = "抽查清單"
SHEET_SUMMARY = "統計摘要"
SHEET_TREND = "月趨勢"
SHEET_UNMATCHED_ICD10 = "ICD10待確認"
SHEET_STAGE_REVIEW = "TNM期別待確認"
SHEET_UNMAPPED_DRUGS = "藥物待確認"
SHEET_UNMAPPED_TREATMENTS = "治療醫令待確認"

VISIBLE_DATA_SHEETS = (
    SHEET_PATIENT_LIST,
    SHEET_TREATMENT_DETAIL,
    SHEET_QA_SAMPLE,
    SHEET_SUMMARY,
    SHEET_TREND,
    SHEET_UNMATCHED_ICD10,
    SHEET_STAGE_REVIEW,
    SHEET_UNMAPPED_DRUGS,
    SHEET_UNMAPPED_TREATMENTS,
)
REVIEW_SHEETS = {
    SHEET_UNMATCHED_ICD10,
    SHEET_STAGE_REVIEW,
    SHEET_UNMAPPED_DRUGS,
    SHEET_UNMAPPED_TREATMENTS,
}
TECHNICAL_SHEETS = (
    "cleaned_data",
    "standardized_data",
    "patient_master",
    "patient_list",
    "stage_review",
)

NAVY = "17365D"
TEAL = "0F6B68"
BLUE = "2F75B5"
GOLD = "BF9000"
LIGHT_BLUE = "DDEBF7"
LIGHT_GOLD = "FFF2CC"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
TEXT = "1F2937"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def format_report_workbook(workbook: Workbook, metrics: Mapping[str, int]) -> None:
    if SHEET_GUIDE in workbook.sheetnames:
        del workbook[SHEET_GUIDE]
    guide = workbook.create_sheet(SHEET_GUIDE, 0)
    _write_guide_sheet(guide, metrics)

    for worksheet in workbook.worksheets:
        if worksheet.title == SHEET_GUIDE:
            continue
        _format_data_sheet(worksheet)

    for sheet_name in TECHNICAL_SHEETS:
        if sheet_name in workbook.sheetnames:
            workbook[sheet_name].sheet_state = "hidden"

    workbook.active = 0


def _write_guide_sheet(worksheet: Worksheet, metrics: Mapping[str, int]) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = TEAL
    worksheet.freeze_panes = "A4"
    worksheet.merge_cells("A1:H1")
    title = worksheet["A1"]
    title.value = "癌症報表清洗結果"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Microsoft JhengHei", size=18, bold=True, color=WHITE)
    title.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 34

    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = f"產生時間：{datetime.now():%Y-%m-%d %H:%M}"
    worksheet["A2"].font = Font(name="Microsoft JhengHei", size=10, color="666666")

    worksheet.merge_cells("A4:H4")
    worksheet["A4"] = "從「病人總表」開始閱讀"
    worksheet["A4"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    worksheet["A4"].font = Font(name="Microsoft JhengHei", size=14, bold=True, color=NAVY)
    worksheet["A4"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[4].height = 28

    worksheet.merge_cells("A5:H5")
    worksheet["A5"] = "病人總表是日常查閱入口；黃色分頁代表需要人工確認。"
    worksheet["A5"].font = Font(name="Microsoft JhengHei", size=11, color=TEXT)

    worksheet["A7"] = "本次結果"
    worksheet["A7"].font = Font(name="Microsoft JhengHei", size=12, bold=True, color=NAVY)
    metric_rows = [
        ("總表筆數", "patient_list_rows"),
        ("不同病歷號數", "unique_mrn_count"),
        ("同病歷號出現多列", "repeated_mrn_count"),
        ("標準化明細筆數", "standardized_rows"),
        ("治療藥物明細筆數", "treatment_detail_rows"),
        ("抽查清單筆數", "qa_sample_rows"),
        ("ICD10 待確認", "unmatched_icd10_rows"),
        ("TNM／期別待確認", "stage_review_rows"),
        ("未對應藥物", "unmapped_drug_rows"),
        ("未對應治療醫令", "unmapped_treatment_rows"),
    ]
    for row_index, (label, key) in enumerate(metric_rows, start=8):
        worksheet.cell(row=row_index, column=1, value=label)
        worksheet.cell(row=row_index, column=2, value=int(metrics.get(key, 0)))
        worksheet.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        worksheet.cell(row=row_index, column=1).font = Font(name="Microsoft JhengHei", bold=True, color=TEXT)
        worksheet.cell(row=row_index, column=2).font = Font(name="Microsoft JhengHei", bold=True, color=BLUE)
        worksheet.cell(row=row_index, column=2).number_format = "#,##0"
        for column in (1, 2):
            worksheet.cell(row=row_index, column=column).border = Border(bottom=THIN_GRAY)

    start_row = 20
    worksheet.cell(row=start_row, column=1, value="工作表")
    worksheet.cell(row=start_row, column=2, value="用途")
    worksheet.cell(row=start_row, column=3, value="何時使用")
    for cell in worksheet[start_row][:3]:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Microsoft JhengHei", bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")

    sheet_guide = [
        (SHEET_PATIENT_LIST, "主要總表，常用欄位在前方", "平常查病人、癌別、期別與治療"),
        (SHEET_TREATMENT_DETAIL, "每筆藥物、手術、放療與其他治療紀錄", "查治療日期、醫令碼與藥物名稱"),
        (SHEET_QA_SAMPLE, "自動挑選 30 位高風險或高資訊量病人", "與院內系統逐筆核對"),
        (SHEET_SUMMARY, "依癌別、期別、性別、年齡與治療彙總", "看整體分布"),
        (SHEET_TREND, "依診斷月份整理趨勢", "看月份變化"),
        (SHEET_UNMATCHED_ICD10, "尚未對應癌別的 ICD10", "更新 ICD10 mapping 前"),
        (SHEET_STAGE_REVIEW, "TNM 與最終期別需確認的組合", "醫療端確認期別時"),
        (SHEET_UNMAPPED_DRUGS, "尚未分類的新藥物醫令", "更新藥物 mapping 前"),
        (SHEET_UNMAPPED_TREATMENTS, "尚未分類的手術／放療等醫令", "更新治療 mapping 前"),
    ]
    for row_index, (sheet_name, purpose, timing) in enumerate(sheet_guide, start=start_row + 1):
        worksheet.cell(row=row_index, column=1, value=sheet_name)
        worksheet.cell(row=row_index, column=1).hyperlink = f"#'{sheet_name}'!A1"
        worksheet.cell(row=row_index, column=1).style = "Hyperlink"
        worksheet.cell(row=row_index, column=2, value=purpose)
        worksheet.cell(row=row_index, column=3, value=timing)
        if sheet_name in REVIEW_SHEETS:
            for column in range(1, 4):
                worksheet.cell(row=row_index, column=column).fill = PatternFill(
                    "solid", fgColor=LIGHT_GOLD
                )

    note_row = start_row + len(sheet_guide) + 3
    worksheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    worksheet.cell(row=note_row, column=1, value="欄位判讀")
    worksheet.cell(row=note_row, column=1).font = Font(
        name="Microsoft JhengHei", size=12, bold=True, color=NAVY
    )
    notes = [
        "「有」表示來源資料中找到該治療紀錄。",
        "治療欄空白表示來源資料未找到紀錄，不等於確認未接受治療。",
        "病人總表以病歷號優先合併；多癌別、多個案與多次治療會保留在摘要及明細。",
        "此檔為資料整理結果，仍需由醫療端確認，不應單獨作為診療判定。",
    ]
    for offset, note in enumerate(notes, start=1):
        worksheet.merge_cells(
            start_row=note_row + offset,
            start_column=1,
            end_row=note_row + offset,
            end_column=8,
        )
        worksheet.cell(row=note_row + offset, column=1, value=f"• {note}")
        worksheet.cell(row=note_row + offset, column=1).font = Font(
            name="Microsoft JhengHei", size=10, color=TEXT
        )

    widths = {"A": 24, "B": 34, "C": 32, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.auto_filter.ref = f"A{start_row}:C{start_row + len(sheet_guide)}"


def _format_data_sheet(worksheet: Worksheet) -> None:
    is_review = worksheet.title in REVIEW_SHEETS
    is_technical = worksheet.title in TECHNICAL_SHEETS
    header_color = GOLD if is_review else (NAVY if is_technical else TEAL)

    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 85 if worksheet.max_column > 18 else 100
    worksheet.freeze_panes = "E2" if worksheet.title in {SHEET_PATIENT_LIST, SHEET_TREATMENT_DETAIL, SHEET_QA_SAMPLE} else "A2"
    worksheet.sheet_properties.tabColor = GOLD if is_review else ("7F8C8D" if is_technical else BLUE)
    worksheet.row_dimensions[1].height = 32

    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(name="Microsoft JhengHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=header_color))

    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        last_column = get_column_letter(worksheet.max_column)
        worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"

    if worksheet.title == SHEET_QA_SAMPLE and worksheet.max_row >= 2:
        validation = DataValidation(
            type="list", formula1='"正確,需修正,待確認"', allow_blank=True
        )
        validation.error = "請選擇：正確、需修正或待確認"
        validation.errorTitle = "抽查結果"
        worksheet.add_data_validation(validation)
        validation.add(f"M2:M{worksheet.max_row}")

    for column_index in range(1, worksheet.max_column + 1):
        header = str(worksheet.cell(row=1, column=column_index).value or "")
        sample_length = len(header)
        for row_index in range(2, min(worksheet.max_row, 80) + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            if value is not None:
                sample_length = max(sample_length, len(str(value)))
        width = min(max(sample_length + 2, 8), 28)
        if "日期" in header or header.endswith("_date") or header in {"診斷日", "生日", "最後回診日"}:
            width = max(width, 13)
        if any(keyword in header for keyword in ("名稱", "摘要", "來源報表", "診斷範例", "備註")):
            width = min(max(width, 18), 32)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
