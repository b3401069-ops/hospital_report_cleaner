from __future__ import annotations

import math
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SEARCH_SHEET = "病人查詢"
PATIENT_SHEET = "病人總表"
TREATMENT_HEADERS = {
    "手術", "放射治療", "化學治療", "標靶治療", "免疫治療", "抗荷爾蒙治療",
    "肝動脈栓塞（TACE）", "射頻燒灼（RFA）",
}
FONT = "Microsoft JhengHei"


def _text_width(value: object) -> int:
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(value or ""))


def format_patient_table(sheet) -> None:
    headers = {cell.value: cell.column for cell in sheet[1]}
    key_column = headers.get("查詢索引", sheet.max_column + 1)
    sheet.cell(1, key_column, "查詢索引")
    mrn_letter = get_column_letter(headers["病歷號"])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, key_column, f'="MRN:"&{mrn_letter}{row}')
    headers["查詢索引"] = key_column
    sheet.freeze_panes = "C2"
    sheet.sheet_view.zoomScale = 85
    sheet.row_dimensions[1].height = 44
    widths = {
        "病患姓名": 12, "病歷號": 14, "性別": 6, "年齡": 7, "癌別摘要": 22,
        "最終期別": 12, "最後回診日": 14, "最後資料日期": 14,
        "期別資料狀態": 23,
        "主責醫師": 12, "個管師": 12, "藥物種類數": 9, "資料提醒": 24,
        "診斷日": 14, "診斷名稱": 40, "TNM／期別摘要": 55,
        "治療摘要": 30, "藥物摘要": 65, "癌症個案數": 10,
        "個案分類": 24, "死亡日期": 14,
    }
    for header, index in headers.items():
        sheet.column_dimensions[get_column_letter(index)].width = (
            9 if header in TREATMENT_HEADERS else widths.get(header, 16)
        )
        color = "526677" if index <= 4 else "416B59" if header in TREATMENT_HEADERS else "376E8A"
        sheet.cell(1, index).fill = PatternFill("solid", fgColor=color)

    for row in sheet.iter_rows(min_row=2):
        visible_lines = 1
        for cell in row:
            header = sheet.cell(1, cell.column).value
            cell.font = Font(name=FONT, size=11, color="253340")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="F3F6F8" if cell.row % 2 else "FFFFFF")
            if header in {"病患姓名", "病歷號"}:
                cell.font = Font(name=FONT, size=11, bold=True, color="253340")
                cell.number_format = "@"
            if header in TREATMENT_HEADERS:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if cell.value == "有":
                    cell.fill = PatternFill("solid", fgColor="DDEFE3")
                    cell.font = Font(name=FONT, size=11, bold=True, color="285A3D")
            if (header == "資料提醒" and cell.value) or (header == "期別資料狀態" and "待確認" in str(cell.value)):
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
                cell.font = Font(name=FONT, size=11, color="845716")
            if header in {"年齡", "藥物種類數", "癌症個案數"}:
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="right", vertical="top")
            if cell.column <= headers["資料提醒"]:
                width = sheet.column_dimensions[cell.column_letter].width - 2
                visible_lines = max(visible_lines, math.ceil(_text_width(cell.value) / max(width, 1)))
        sheet.row_dimensions[row[0].row].height = min(100, max(32, visible_lines * 16 + 6))

    # Full summaries remain in the table and in the single-patient view.
    first_detail = headers["診斷日"]
    for index in range(first_detail, sheet.max_column + 1):
        dimension = sheet.column_dimensions[get_column_letter(index)]
        dimension.outlineLevel = 1
        dimension.hidden = True
    sheet.column_dimensions[get_column_letter(sheet.max_column + 1)].collapsed = True
    sheet.sheet_properties.outlinePr.summaryRight = True
    if sheet.max_row > 1:
        table = Table(displayName="PatientOverview", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        sheet.add_table(table)
    sheet.print_title_rows = "1:1"


def write_patient_search(workbook) -> None:
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    if SEARCH_SHEET in workbook.sheetnames:
        del workbook[SEARCH_SHEET]
    sheet = workbook.create_sheet(SEARCH_SHEET, 1)
    patients = workbook[PATIENT_SHEET]
    headers = {cell.value: cell.column for cell in patients[1]}
    last_row = max(patients.max_row, 3)
    mrn_column = get_column_letter(headers["查詢索引"])
    mrns = f"'{PATIENT_SHEET}'!${mrn_column}$2:${mrn_column}${last_row}"
    # A text prefix prevents numeric coercion; escape literal wildcard characters.
    key = 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE("MRN:"&$B$3,"~","~~"),"*","~*"),"?","~?")'
    sheet["H2"] = f'=IF($B$3="",0,COUNTIF({mrns},"="&{key}))'
    sheet["H3"] = f'=IF($H$2=1,MATCH({key},{mrns},0),0)'
    sheet.column_dimensions["H"].hidden = True
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "376E8A"
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 100
    sheet["A1"] = "病人查詢"
    sheet["A1"].font = Font(name=FONT, size=18, bold=True, color="253340")
    sheet.row_dimensions[1].height = 32
    sheet["B2"] = "輸入病歷號查單人；比較多人請開啟「病人總表」。"
    sheet["B2"].font = Font(name=FONT, size=10, color="61717E")
    sheet["A3"] = "病歷號"
    sheet["B3"].number_format = "@"
    sheet["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["B3"].font = Font(name=FONT, size=14, bold=True, color="174B76")
    sheet["B3"].border = Border(bottom=Side(style="medium", color="C89B38"))
    sheet.row_dimensions[3].height = 30
    sheet["A4"] = "查詢結果"
    sheet["B4"] = '=IF(B3="","請輸入完整病歷號，保留前面的 0",IF(H2=0,"查無此病歷號",IF(H2>1,"同病歷號有多列，請至病人總表核對","已找到病人")))'

    fields = [
        (6, "病患姓名"), (7, "性別"), (8, "年齡"), (9, "癌別摘要"),
        (10, "最終期別"), (11, "診斷日"), (12, "診斷名稱"),
        (13, "期別資料狀態"),
        (14, "TNM／期別摘要"), (16, "治療摘要"), (17, "藥物種類數"),
        (18, "藥物摘要"), (20, "最後回診日"), (21, "最後資料日期"),
        (22, "主責醫師"), (23, "個管師"), (24, "癌症個案數"),
        (25, "個案分類"), (26, "死亡日期"), (27, "資料提醒"),
    ]
    for row, label in fields:
        sheet.cell(row, 1, label)
        if label not in headers:
            continue
        column = get_column_letter(headers[label])
        reference = f"INDEX('{PATIENT_SHEET}'!${column}$2:${column}${last_row},$H$3,1)"
        sheet.cell(row, 2, f'=IF($H$3=0,"",IF({reference}="","",{reference}))')
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="EDF2F5")
        sheet.cell(row, 1).font = Font(name=FONT, size=11, color="4D6374")
        sheet.cell(row, 2).font = Font(name=FONT, size=11, color="253340")
        for cell in sheet[row][:2]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="DFE5EA"))
        max_lines = max(
            (sum(max(1, math.ceil(_text_width(line) / 95)) for line in str(values[0] or "").splitlines())
             for values in patients.iter_rows(min_row=2, min_col=headers[label], max_col=headers[label], values_only=True)),
            default=1,
        )
        sheet.row_dimensions[row].height = min(409, max(25, max_lines * 16 + 8))
        if label in {"診斷日", "最後回診日", "最後資料日期", "死亡日期"}:
            sheet.cell(row, 2).number_format = "yyyy-mm-dd"
    sheet["A29"] = "病人總表"
    sheet["B29"] = f'=IF(H3=0,"",HYPERLINK("#\'{PATIENT_SHEET}\'!A"&(H3+1),"前往這位病人的總表列"))'
    sheet["A30"] = "治療藥物明細"
    sheet["B30"] = "開啟明細後，以病歷號欄篩選此病人"
    sheet["B30"].hyperlink = "#'治療藥物明細'!A1"
    sheet["B30"].style = "Hyperlink"
    sheet["A32"] = "有＝找到治療紀錄；空白＝未提供。多癌別或多期別保留原摘要。"
    sheet["A32"].font = Font(name=FONT, size=10, color="61717E")
    sheet.print_options.horizontalCentered = True
    sheet.print_area = "A1:B30"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
