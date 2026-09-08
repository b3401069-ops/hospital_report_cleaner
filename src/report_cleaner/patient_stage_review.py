from __future__ import annotations

from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import shutil

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


SHEET = "病人期別待確認"
FILE_NAME = "病人期別確認.xlsx"
IDENTITY = {
    "patient_key": "病人識別鍵", "medical_record_no": "病歷號", "case_id": "個案編號",
    "cancer_type": "癌別", "diagnosis_date": "診斷日",
    "clinical_tnm": "臨床TNM", "pathologic_tnm": "病理TNM",
}
INPUTS = ["確認期別", "分期類別", "確認依據", "確認人", "確認日期", "確認狀態"]
STAGE_VALUES = set("0 1 2 3 4 I IA IA1 IA2 IA3 IB IB1 IB2 IB3 IC IC1 IC2 IC3 IS II IIA IIA1 IIA2 IIB IIC III IIIA IIIA1 IIIA2 IIIB IIIC IIIC1 IIIC2 IV IVA IVA1 IVA2 IVB IVC 不適用".split())
STAGING_TYPES = {"臨床", "病理", "治療後臨床", "治療後病理", "其他（請註明）"}
COLUMNS = [
    "病患姓名", "病歷號", "個案編號", "癌別", "診斷日", "臨床TNM", "病理TNM",
    "現有期別", "待確認原因", "主責醫師", "個管師", "來源檔案", "來源報表",
    "來源筆數", *INPUTS, "本次狀態", "確認鍵", "病人識別鍵",
]


def text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def confirmation_key(row) -> str:
    values = [text(row.get(column, "")) for column in IDENTITY]
    return hashlib.sha256(json.dumps(values, ensure_ascii=False).encode("utf-8")).hexdigest()


def review_path(root: Path) -> Path:
    return root / "review" / FILE_NAME


def _prepare(frame: pd.DataFrame) -> pd.DataFrame:
    working = frame.copy()
    for column in [*IDENTITY, "final_stage", "patient_name", "doctor", "manager", "source_file", "report_name"]:
        if column not in working:
            working[column] = ""
        working[column] = working[column].map(text)
    working["_confirmation_key"] = [confirmation_key(row) for row in working.to_dict("records")]
    return working


def _read_records(path: Path) -> list[dict]:
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "個案確認" not in workbook.sheetnames:
            raise ValueError("病人期別確認檔缺少「個案確認」工作表，請保留原格式。")
        rows = workbook["個案確認"].iter_rows(values_only=True)
        headers = next(rows, ())
        if not set(COLUMNS).issubset(headers):
            raise ValueError("病人期別確認檔的欄名不完整，請保留原格式。")
        records = [dict(zip(headers, (text(value) for value in row))) for row in rows if any(row)]
        keys = [record["確認鍵"] for record in records]
        if len(keys) != len(set(keys)):
            raise ValueError("病人期別確認檔有重複確認鍵，請勿複製整列。")
        for record in records:
            context = {column: record[label] for column, label in IDENTITY.items()}
            if record["確認鍵"] != confirmation_key(context):
                raise ValueError("病人期別確認檔的病人或 TNM 識別欄被修改，請只填黃色確認欄。")
            if record["確認狀態"] == "已確認":
                if any(not record[column] for column in INPUTS):
                    raise ValueError("標為已確認的個案，必須填寫期別、分期類別、依據、確認人與日期。")
                if not (record["病歷號"] or record["個案編號"]):
                    raise ValueError("個案確認必須有病歷號或個案編號，不能僅以姓名套用。")
                if not record["癌別"] or not (record["個案編號"] or record["診斷日"]):
                    raise ValueError("個案識別不足：請先補正來源的癌別及個案編號或診斷日，重新清洗後再確認。")
                if any(record[column].startswith("=") for column in INPUTS):
                    raise ValueError("個案確認欄請填文字或日期，不可填公式。")
                if record["確認期別"].upper() not in STAGE_VALUES or record["分期類別"] not in STAGING_TYPES:
                    raise ValueError("確認期別或分期類別無法辨識。期別請填如 2、IIA；尚未確定請保持待確認。")
                if pd.isna(pd.to_datetime(record["確認日期"], errors="coerce")):
                    raise ValueError("個案確認日期無法辨識，請填西元年-月-日。")
        return records
    finally:
        workbook.close()


def apply_confirmations(frame: pd.DataFrame, path: Path) -> pd.DataFrame:
    records = _read_records(path)
    confirmed = {record["確認鍵"]: record for record in records if record["確認狀態"] == "已確認"}
    if frame.empty or not confirmed:
        return frame.copy()
    result = frame.copy()
    if "final_stage_source" not in result:
        result["final_stage_source"] = ""
    for index, row in result.iterrows():
        record = confirmed.get(confirmation_key(row))
        if record and not text(row.get("final_stage", "")):
            result.at[index, "final_stage"] = record["確認期別"]
            result.at[index, "final_stage_source"] = "patient_confirmed"
    return result


def build_queue(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=COLUMNS)
    working = _prepare(frame)
    has_tnm = working["clinical_tnm"].ne("") | working["pathologic_tnm"].ne("")
    staging_report = working["report_name"].str.contains("D1\\.2|收案|個案清單|計畫書|治療計畫", regex=True)
    candidates = working.loc[has_tnm | staging_report].copy()
    represented = set(zip(candidates["patient_key"], candidates["cancer_type"]))
    # A drug/order row lacking stage is not automatically a staging defect.
    # Add one fallback only when this patient's cancer has no stage anywhere.
    fallbacks = []
    for key, group in working.groupby(["patient_key", "cancer_type"], sort=False):
        if key not in represented and not group["final_stage"].ne("").any():
            fallbacks.append(group.iloc[[0]])
    if fallbacks:
        candidates = pd.concat([candidates, *fallbacks], ignore_index=True)
    known_keys = set(working.loc[working["final_stage"].ne(""), "_confirmation_key"])
    candidates = candidates.loc[~candidates["_confirmation_key"].isin(known_keys)]
    rows = []
    for key, group in candidates.groupby("_confirmation_key", sort=False):
        first = group.iloc[0]
        clinical, pathologic = first["clinical_tnm"], first["pathologic_tnm"]
        tnm = (pathologic or clinical).upper()
        if not tnm:
            reason = "缺期別及TNM，需查院內紀錄"
        elif tnm in {"CTNM", "PTNM", "TNM"}:
            reason = "TNM內容不具體，缺期別"
        elif "X" in tnm or "?" in tnm:
            reason = "TNM有未知項目，缺期別"
        else:
            reason = "有TNM但缺期別，需核對"
        if not first["cancer_type"] or not (first["case_id"] or first["diagnosis_date"]):
            reason += "；個案識別不足，請先補正來源"
        record = {label: first[column] for column, label in IDENTITY.items()}
        for source, label in [
            ("patient_name", "病患姓名"), ("doctor", "主責醫師"), ("manager", "個管師"),
            ("source_file", "來源檔案"), ("report_name", "來源報表"),
        ]:
            record[label] = " | ".join(dict.fromkeys(value for value in group[source] if value))
        record.update({"確認鍵": key, "待確認原因": reason, "來源筆數": len(group),
                       "現有期別": "", "本次狀態": "尚待確認", "確認狀態": "待確認"})
        rows.append(record)
    return pd.DataFrame(rows, columns=COLUMNS).fillna("")


def status_by_patient(patient_frame: pd.DataFrame, queue: pd.DataFrame) -> pd.Series:
    pending = set(queue["病人識別鍵"]) if not queue.empty else set()
    values = []
    for row in patient_frame.to_dict("records"):
        stage = text(row.get("final_stage", ""))
        if text(row.get("patient_key", "")) in pending:
            values.append("部分資料待確認" if stage else "缺期別待確認")
        elif not stage:
            values.append("缺期別待確認")
        elif "patient_confirmed" in text(row.get("final_stage_source", "")):
            values.append("含個案人工確認")
        else:
            values.append("已有期別（未核驗）")
    return pd.Series(values, index=patient_frame.index, dtype="object")


def sync_review(frame: pd.DataFrame, path: Path) -> Path:
    existing = {record["確認鍵"]: record for record in _read_records(path)}
    current_queue = build_queue(frame)
    current = _prepare(frame)
    stages = current.groupby("_confirmation_key")["final_stage"].agg(lambda values: set(values) - {""}).to_dict()
    for record in current_queue.to_dict("records"):
        key = record["確認鍵"]
        if key in existing:
            record.update({column: existing[key][column] for column in INPUTS})
        existing[key] = record
    for key, record in existing.items():
        current_stages = stages.get(key)
        if current_stages is None:
            state = "本批無相同個案資料，保留確認紀錄"
        elif current_stages:
            state = "本批已有期別"
            if record["確認狀態"] == "已確認":
                state = "已套用或來源一致" if current_stages == {record["確認期別"]} else "來源期別與確認值不同，請核對"
        else:
            state = "尚待確認"
        record["本次狀態"] = state
        record["現有期別"] = " | ".join(sorted(current_stages or []))

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "填寫說明"
    for row in [
        ["病人期別確認"],
        ["在「個案確認」黃色欄位填寫核對結果，完成後選「已確認」，儲存並關閉本檔，再執行清洗報表。"],
        ["只補同一病人、個案、癌別、診斷日及TNM的空白期別；不新增共用規則。"],
        ["已有不同期別時保留來源，並在本次狀態提示核對；本檔不判斷臨床分期是否正確。"],
        ["資料不足請維持待確認。請勿修改灰色識別欄或複製整列。"],
        ["每次執行保留已填內容並補入新個案；舊批次的確認紀錄也會保留。"],
        ["本檔含病人資料。移至其他電腦使用時，請連同 review 資料夾一起移轉。"],
    ]:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 110
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row[0].row].height = 38
    sheet = workbook.create_sheet("個案確認")
    sheet.append(COLUMNS)
    for record in sorted(existing.values(), key=lambda row: (row["確認狀態"] == "已確認", row["病歷號"], row["癌別"], row["確認鍵"])):
        sheet.append([record.get(column, "") for column in COLUMNS])
    format_queue_sheet(sheet, editable=True)
    workbook.active = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    try:
        workbook.save(temporary)
        if path.exists():
            shutil.copy2(path, path.with_name(path.stem + ".backup.xlsx"))
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return path


def format_queue_sheet(sheet, *, editable: bool = False) -> None:
    from openpyxl.utils import get_column_letter

    sheet.freeze_panes = "D2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "BF9000"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(COLUMNS, 1):
        column = sheet.column_dimensions[get_column_letter(index)]
        column.width = 32 if header in {"待確認原因", "確認依據", "來源檔案", "來源報表", "本次狀態"} else 17
        if header in {"確認鍵", "病人識別鍵"} or (not editable and header in INPUTS):
            column.hidden = True
        for cells in sheet.iter_rows(min_col=index, max_col=index):
            cell = cells[0]
            cell.font = Font(name="Microsoft JhengHei", size=11, bold=cell.row == 1,
                             color="FFFFFF" if cell.row == 1 else "253340")
            cell.fill = PatternFill("solid", fgColor="9A751A" if cell.row == 1 else "FFF2CC" if editable and header in INPUTS else "F3F6F8")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if header in {"病歷號", "個案編號", "確認期別"}:
                cell.number_format = "@"
    for index in range(1, sheet.max_row + 1):
        sheet.row_dimensions[index].height = 48 if index > 1 else 32
    if editable:
        for header, choices in [("確認狀態", '"待確認,已確認"'), ("分期類別", '"臨床,病理,治療後臨床,治療後病理,其他（請註明）"')]:
            validation = DataValidation(type="list", formula1=choices, allow_blank=True)
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            letter = get_column_letter(COLUMNS.index(header) + 1)
            validation.add(f"{letter}2:{letter}{max(sheet.max_row, 2)}")
